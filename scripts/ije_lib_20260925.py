# -*- coding: utf-8 -*-
"""
ije_lib_20260925.py

Infra da versao IJE:
  - rastreador de numeros (cada valor impresso registra arquivo, filtro e coluna de origem)
  - referencias numeradas por ordem de aparicao (estilo Oxford SCIMED/Vancouver, [1], [2])
  - utilitarios de Word (destaque amarelo para campos que dependem de confirmacao humana)
Nenhuma estimativa nova e' calculada aqui: so' leitura e formatacao dos CSVs existentes,
com as MESMAS funcoes de arredondamento do manuscrito SSM (ms_numeros_20260925).
"""
import re, html, sys
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, str(Path(__file__).parent))
from ms_numeros_20260925 import *            # noqa: F401,F403  (t1..m3, aten, reg, boot, dm, f1,f2,f3,n0,ci,pfmt, T1.., AUC, N_STRATA, BASE)
from refs_ssm_20260925 import DOIS, ANO, SUFIXO, carregar



def n0(x):
    """Oxford SciMed: separador de milhar = espaco (nao separavel) so' a partir de 10 000"""
    v = int(round(float(x)))
    return f"{v:,}".replace(",", "\u00a0") if abs(v) >= 10000 else str(v)


def ci(lo, hi, dec=2):
    fmt = {1: f1, 2: f2, 3: f3}[dec]
    return f"{fmt(lo)} to {fmt(hi)}" if lo < 0 else f"{fmt(lo)}-{fmt(hi)}"


OXF_RE = re.compile(r"\b(characteris|racialis|dichotomis|linearis|criticis|recognis|anonymis|normalis|standardis|summaris|organis|minimis|prioritis|marginalis|contextualis|harmonis|categoris|utilis|realis|optimis|randomis|stabilis|visualis|operationalis|conceptualis|generalis|specialis|finalis|maximis|socialis|hospitalis|centralis|authoris|sensitis|initialis|customis|formalis|italicis|capitalis)(e|ed|es|ing|ation|ations)\b", re.I)
OXF_RE2 = re.compile(r"\bgeneralisability\b", re.I)


def oxford(t):
    """ortografia de Oxford (IJE mini checklist): -ize, mas 'analyse'. Regex por palavra inteira: NAO altera 'characteristic(s)'"""
    def sub(m):
        stem, suf = m.group(1), m.group(2)
        new = stem[:-1] + "z"
        out = new + suf
        return out.capitalize() if m.group(0)[0].isupper() else out
    t = OXF_RE.sub(sub, t)
    return OXF_RE2.sub(lambda m: "Generalizability" if m.group(0)[0].isupper() else "generalizability", t)


# numeracao unica dos materiais suplementares (texto principal e suplemento usam esta tabela)
ST = {"coding": "S1", "table1": "S2", "own": "S3", "pairs": "S4", "cells": "S5", "joint": "S6", "three": "S7", "region": "S8", "strata": "S9",
      "grad": "S10", "maihda": "S11", "fixed": "S12", "weights": "S13", "var": "S15", "comp": "S16", "std": "S17", "cut": "S18", "ign": "S19", "samp": "S20", "strobe": "S21"}


OUT = BASE / "MANUSCRITO_IJE_20260925"; OUT.mkdir(exist_ok=True)
FIG = OUT / "figures"
META = carregar()

# ============================================================== rastreador
TR = []; CUR = {"loc": ""}
SSM = {"t1": "Table 1; §3.1", "t2": "Table 3 & Fig. 2 (§3.3); Suppl. Table S2 (cell prevalences, decomposition)",
       "t3": "Table 4 (§3.4)", "t4": "§3.6, Fig. 3; Suppl. Table S6a", "t5": "§3.6; Suppl. Table S6b",
       "t6": "§3.5 & Fig. 1; Suppl. Table S5", "t7": "§3.4; Suppl. Table S4", "aten": "Table 2 (§3.2)",
       "m1": "Table 5 (§3.7); Suppl. Table S8", "m2": "Suppl. Table S7", "m3": "Fig. 4; Supplementary_data_strata.xlsx",
       "auc": "Table 5 (§3.7)", "jd": "§3.3 & Suppl. Table S2 (decomposition of the joint prevalence difference)",
       "reg": "Suppl. Table S3", "abs": "Abstract (rewritten for IJE; all numbers below are from the sources listed)"}


def setloc(s): CUR["loc"] = s


def log(item, printed, raw, src, ssm, flag=""):
    TR.append({"location": CUR["loc"], "item": item, "printed": printed, "raw_value": raw, "source": src, "ssm_manuscript_location": ssm, "flag": flag})


def _num(x):
    try: return float(x)
    except Exception: return x


def _ciw(mode, est, lo, hi, dec):
    fmt = {1: f1, 2: f2, 3: f3}[dec]
    if mode == "tab": return f"{fmt(est)} ({ci(lo, hi, dec)})"
    if mode == "sq": return f"{fmt(est)} [{ci(lo, hi, dec)}]"
    if mode == "br": return f"{fmt(est)} [95% CI {ci(lo, hi, dec)}]"
    return f"{fmt(est)} (95% CI {ci(lo, hi, dec)})"


def _src(fname, sel, col):
    if fname.startswith("rodada_D/"): return f"MANUSCRITO_IJE_20260925/{fname} | {sel} | {col}"
    return f"dados/{DVER}/{fname} | {sel} | {col}"


# ---- t1 (descritiva)
def q1(row, col, fmt=f1):
    raw = row[col]; s = fmt(raw)
    log(f"{row.variavel}/{row.categoria}: {col}", s, _num(raw), _src("t1_descritiva.csv", f"variavel={row.variavel}; categoria={row.categoria}", col), SSM["t1"]); return s


def q1ci(row, est, lo, hi, dec=1, mode="text"):
    s = _ciw(mode, row[est], row[lo], row[hi], dec)
    log(f"{row.variavel}/{row.categoria}: {est} (CI {lo},{hi})", s, f"{_num(row[est])}|{_num(row[lo])}|{_num(row[hi])}", _src("t1_descritiva.csv", f"variavel={row.variavel}; categoria={row.categoria}", f"{est},{lo},{hi}"), SSM["t1"]); return s


def q1pc(row, est, lo, hi, dec=1, short=False):
    """prevalencia em % com IC: '27.6% (95% CI 27.1-28.1)' (short: sem '95% CI')"""
    s = f"{f1(row[est])}% ({'' if short else '95% CI '}{ci(row[lo], row[hi], dec)})"
    log(f"{row.variavel}/{row.categoria}: {est} (CI)", s, f"{_num(row[est])}|{_num(row[lo])}|{_num(row[hi])}", _src("t1_descritiva.csv", f"variavel={row.variavel}; categoria={row.categoria}", f"{est},{lo},{hi}"), SSM["t1"]); return s


# ---- t2 (pares)
def q2(row, col, fmt=f2):
    raw = row[col]; s = fmt(raw)
    log(f"{row.par} / {row.desfecho}: {col}", s, _num(raw), _src("t2_pares.csv", f"par={row.par}; desfecho={row.desfecho}", col), SSM["t2"]); return s


def q2ci(row, est, lo, hi, dec=2, mode="text"):
    s = _ciw(mode, row[est], row[lo], row[hi], dec)
    log(f"{row.par} / {row.desfecho}: {est} (CI)", s, f"{_num(row[est])}|{_num(row[lo])}|{_num(row[hi])}", _src("t2_pares.csv", f"par={row.par}; desfecho={row.desfecho}", f"{est},{lo},{hi}"), SSM["t2"]); return s


def qjdci(row, idx, dec=1):
    """componente absoluto da decomposicao (pp) com IC de 95% da configuracao C (colunas de t2 v3): idx 0 conjunta, 1 raca, 2 SES, 3 interacao"""
    col = ["conjunta_pp", "comp_raca_pp", "comp_ses_pp", "interacao_pp"][idx]; fmt = {1: f1, 2: f2}[dec]
    s = f"{fmt(row[col])} (95% CI {ci(row[col + '_lo'], row[col + '_hi'], dec)})"
    log(f"{row.par} / {row.desfecho}: {col} (CI, config. C)", s, f"{_num(row[col])}|{_num(row[col + '_lo'])}|{_num(row[col + '_hi'])}", _src("t2_pares.csv", f"par={row.par}; desfecho={row.desfecho}", f"{col},{col}_lo,{col}_hi"), SSM["jd"]); return s


def qshare(row, with_ci=False):
    """participacao percentual da interacao na diferenca conjunta (estimativa descritiva; IC so' no suplemento)"""
    col = "participacao_pct"; s = f"{f1(row[col])}" + (f" (95% CI {ci(row[col + '_lo'], row[col + '_hi'], 1)})" if with_ci else "")
    log(f"{row.par} / {row.desfecho}: {col}", s, f"{_num(row[col])}|{_num(row[col + '_lo'])}|{_num(row[col + '_hi'])}", _src("t2_pares.csv", f"par={row.par}; desfecho={row.desfecho}", f"{col}" + (",lo,hi" if with_ci else "")), SSM["jd"]); return s


def qstd(y, q, which="padronizado", dec=2, with_ci=True, mode="text"):
    """D-1: analise padronizada (sensibilidade descritiva) vs bruta; which = padronizado | bruto"""
    r = rd_std[(rd_std.desfecho == y) & (rd_std.quantidade == q)].iloc[0]; e, lo, hi = r[which + "_est"], r[which + "_lo"], r[which + "_hi"]
    s = _ciw(mode, e, lo, hi, dec) if with_ci else {1: f1, 2: f2}[dec](e)
    log(f"standardized sensitivity ({which}) {y}: {q}", s, f"{_num(e)}|{_num(lo)}|{_num(hi)}", _src("rodada_D/D1_padronizado_vs_bruto_raca_renda.csv", f"desfecho={y}; quantidade={q}", f"{which}_est,{which}_lo,{which}_hi"), "rodada D (D-1)"); return s


def qcut(cut, y, q, dec=2, mode="text", with_ci=True):
    """D-3: sensibilidade ao ponto de corte; cut em {'<=1/4 SM (primario)', '<=1/2 SM', '<=1 SM'}; config. C"""
    r = rd_cut[(rd_cut.corte == cut) & (rd_cut.desfecho == y) & (rd_cut.quantidade == q) & (rd_cut.config == "C")].iloc[0]
    s = _ciw(mode, r.estimativa, r.ic_inf, r.ic_sup, dec) if with_ci else {1: f1, 2: f2}[dec](r.estimativa)
    log(f"cut-point {cut} {y}: {q}", s, f"{_num(r.estimativa)}|{_num(r.ic_inf)}|{_num(r.ic_sup)}", _src("rodada_D/D3_sensibilidade_corte_renda.csv", f"corte={cut}; desfecho={y}; quantidade={q}; config=C", "estimativa,ic_inf,ic_sup"), "rodada D (D-3)"); return s


def qign(classe, analise="bruta", dec=3):
    """D-5: maior diferenca absoluta por classe de estimando ao excluir os registros ignorados"""
    r = rd_ign_sum[(rd_ign_sum.classe == classe) & (rd_ign_sum.analise == analise)].iloc[0]; v = abs(r.maior_dif_abs); s = f"{v:.{dec}f}"
    log(f"ignored-records exclusion: max abs difference, {classe} ({analise})", s, float(r.maior_dif_abs), _src("rodada_D/D5_maiores_diferencas_por_classe.csv", f"classe={classe}; analise={analise}", "maior_dif_abs"), "rodada D (D-5)"); return s


def qign_max(classes, analise="bruta", dec=3):
    r = rd_ign_sum[(rd_ign_sum.classe.isin(classes)) & (rd_ign_sum.analise == analise)]; v = float(r.maior_dif_abs.abs().max()); s = f"{v:.{dec}f}"
    log(f"ignored-records exclusion: max abs difference over {'; '.join(classes)} ({analise})", s, v, _src("rodada_D/D5_maiores_diferencas_por_classe.csv", f"classe in {classes}; analise={analise}", "maior_dif_abs (max of absolute values)"), "rodada D (D-5)"); return s


def qign_q5(y="ia_total", dec=2):
    r = rd_ign[(rd_ign.grupo == "prevalencia") & rd_ign.par.str.startswith("renda q5") & (rd_ign.desfecho == y)].iloc[0]; s = f"{abs(r.dif_abs):.{dec}f}"
    log(f"ignored-records exclusion: >2 MW band prevalence change ({y})", s, float(r.dif_abs), _src("rodada_D/D5_ignorados_comparacao_completa.csv", f"grupo=prevalencia; par=renda q5; desfecho={y}", "dif_abs"), "rodada D (D-5)"); return s


def qjd(row, idx, key=None):
    """decomposicao da diferenca de prevalencia conjunta (pp), derivada das 4 prevalencias ponderadas de t2 (mesma formula do manuscrito SSM)"""
    joint = row.p11 - row.p00; a = row.p10 - row.p00; b = row.p01 - row.p00; it = row.p11 - row.p10 - row.p01 + row.p00
    val = [joint, a, b, it, it / joint * 100][idx]
    nm = ["joint difference (p11-p00)", "race alone (p10-p00)", "SES alone (p01-p00)", "interaction contrast (p11-p10-p01+p00)", "interaction share of joint difference (%)"][idx]
    s = f1(val)
    log(f"{row.par} / {row.desfecho}: {nm}", s, float(val), _src("t2_pares.csv", f"par={row.par}; desfecho={row.desfecho}", "p00,p10,p01,p11 (formula in Methods)"), SSM["jd"],
        "derived from the four weighted cell prevalences (same formula as SSM manuscript); no CI exists")
    return s


# ---- t3, t6, aten, m1, m2, auc
def q3ci(termo, desf, mode="text", dec=2):
    r = t3[(t3.termo == termo) & (t3.desfecho == desf)].iloc[0]; s = _ciw(mode, r.rr, r.lo, r.hi, dec)
    log(f"joint model {termo} / {desf}", s, f"{r.rr}|{r.lo}|{r.hi}", _src("t3_conjunto.csv", f"termo={termo}; desfecho={desf}", "rr,lo,hi"), SSM["t3"]); return s


def q6(expo, por, estrato, desf, col, fmt=f2):
    r = T6(expo, por, estrato, desf); raw = r[col]; s = fmt(raw)
    log(f"{expo} by {por}={estrato} / {desf}: {col}", s, _num(raw), _src("t6_estratificado.csv", f"exposicao={expo}; estratificado_por={por}; estrato={estrato}; desfecho={desf}", col), SSM["t6"]); return s


def q6ci(expo, por, estrato, desf, est, lo, hi, dec=2, mode="text"):
    r = T6(expo, por, estrato, desf); s = _ciw(mode, r[est], r[lo], r[hi], dec)
    log(f"{expo} by {por}={estrato} / {desf}: {est} (CI)", s, f"{_num(r[est])}|{_num(r[lo])}|{_num(r[hi])}", _src("t6_estratificado.csv", f"exposicao={expo}; estratificado_por={por}; estrato={estrato}; desfecho={desf}", f"{est},{lo},{hi}"), SSM["t6"]); return s


def qa(var, esp, desf, col, fmt=f2):
    r = ATEN(var, esp, desf); raw = r[col]; s = fmt(raw)
    log(f"attenuation: {var}/{esp}/{desf}: {col}", s, _num(raw), _src("../atenuacao_renda_raca_sexo_2023.csv", f"variavel={var}; especificacao={esp}; desfecho={desf}", col), SSM["aten"]); return s


def qaci(var, esp, desf, mode="text", dec=2):
    r = ATEN(var, esp, desf); s = _ciw(mode, r.rr, r.ic95_low, r.ic95_high, dec)
    log(f"attenuation: {var}/{esp}/{desf}: rr (CI)", s, f"{r.rr}|{r.ic95_low}|{r.ic95_high}", _src("../atenuacao_renda_raca_sexo_2023.csv", f"variavel={var}; especificacao={esp}; desfecho={desf}", "rr,ic95_low,ic95_high"), SSM["aten"]); return s


def qm1(desf, col, fmt=f1, scale=1.0):
    raw = m1.loc[desf, col]; s = fmt(raw * scale)
    log(f"MAIHDA {desf}: {col}", s, _num(raw), _src("m1_maihda_geral.csv", f"desfecho={desf}", col), SSM["m1"]); return s


def qm1ci(desf, est, lo, hi, dec=1, scale=100.0, mode="text"):
    s = _ciw(mode, m1.loc[desf, est] * scale, m1.loc[desf, lo] * scale, m1.loc[desf, hi] * scale, dec)
    log(f"MAIHDA {desf}: {est} (CI)", s, f"{m1.loc[desf, est]}|{m1.loc[desf, lo]}|{m1.loc[desf, hi]}", _src("m1_maihda_geral.csv", f"desfecho={desf}", f"{est},{lo},{hi} (x100)"), SSM["m1"]); return s


def qm1pc(desf, est, lo, hi, dec=1, short=False):
    """percentual MAIHDA com IC: '21.5% (95% CI 18.1-25.8)'"""
    e, l_, h_ = m1.loc[desf, est] * 100, m1.loc[desf, lo] * 100, m1.loc[desf, hi] * 100
    s = f"{f1(e)}% ({'' if short else '95% CI '}{ci(l_, h_, dec)})"
    log(f"MAIHDA {desf}: {est} (CI)", s, f"{m1.loc[desf, est]}|{m1.loc[desf, lo]}|{m1.loc[desf, hi]}", _src("m1_maihda_geral.csv", f"desfecho={desf}", f"{est},{lo},{hi} (x100)"), SSM["m1"]); return s


def qauc(desf, key, fmt=f2):
    raw = AUC[desf][key]; s = fmt(raw)
    log(f"AUC {desf}: {key}", s, float(raw), "computed in ms_numeros_20260925.calcula_auc() from m3_maihda_estratos.csv + dados_maihda_2023.csv", SSM["auc"]); return s


# ============================================================== referencias numeradas
MANUAL = {
    "crenshaw1989": "Crenshaw K. Demarginalizing the intersection of race and sex: a Black feminist critique of antidiscrimination doctrine, feminist theory and antiracist politics. *Univ Chic Leg Forum* 1989;1989:139-67.",
    "fao2025": "FAO, IFAD, UNICEF, WFP, WHO. The State of Food Security and Nutrition in the World 2025. Rome: FAO, 2025. doi:10.4060/cd6008en",
    "ibge2024": "Instituto Brasileiro de Geografia e Estatística. Pesquisa Nacional por Amostra de Domicílios Contínua: Segurança Alimentar 2023. Rio de Janeiro: IBGE, 2024.",
    "camara2026": "Câmara JHR, de Castro Junior PCP, da Silva Luiz GB et al. Racial and territorial inequalities in food insecurity in one of the largest metropolises in Latin America: Rio de Janeiro, Brazil, 2023-2024. *Am J Public Health* 2026; published online 23 July 2026:e1-e10. doi:10.2105/AJPH.2026.308566 ⟦CONFIRM against the publisher page: metadata taken from the PubMed record (PMID 42492040); volume and issue not yet assigned⟧",
}
ABREV = {"Social Science & Medicine": "Soc Sci Med", "SSM - Population Health": "SSM Popul Health", "American Journal of Epidemiology": "Am J Epidemiol",
         "International Journal of Epidemiology": "Int J Epidemiol", "Epidemiologic Methods": "Epidemiol Methods", "European Journal of Epidemiology": "Eur J Epidemiol",
         "Social Psychiatry and Psychiatric Epidemiology": "Soc Psychiatry Psychiatr Epidemiol", "Advances in Nutrition": "Adv Nutr",
         "Cadernos de Saúde Pública": "Cad Saude Publica", "PLOS Global Public Health": "PLOS Glob Public Health", "Annual Review of Sociology": "Annu Rev Sociol",
         "Annual Review of Public Health": "Annu Rev Public Health", "The Lancet": "Lancet", "Saúde e Sociedade": "Saude Soc",
         "Journal of Health and Social Behavior": "J Health Soc Behav", "Social Issues and Policy Review": "Soc Issues Policy Rev", "Health Affairs": "Health Aff",
         "The Journal of Nutrition": "J Nutr", "Food and Nutrition Bulletin": "Food Nutr Bull", "Applied Economic Perspectives and Policy": "Appl Econ Perspect Policy",
         "Canadian Journal of Public Health": "Can J Public Health", "American Journal of Public Health": "Am J Public Health",
         "Paediatric and Perinatal Epidemiology": "Paediatr Perinat Epidemiol", "Research Integrity and Peer Review": "Res Integr Peer Rev",
         "Stanford Law Review": "Stanford Law Rev", "Food Policy": "Food Policy", "BMC Public Health": "BMC Public Health"}
TITLE_OVR = {"rothman1980": "Concepts of interaction",
             "cezimbra2026a": "Does the Brazilian Household Food Insecurity Measurement Scale allow for valid comparisons across sex, race, and education? A measurement invariance analysis using the National Household Sample Survey",
             "perezescamilla2004": "An adapted version of the U.S. Department of Agriculture Food Insecurity Module is a valid tool for assessing household food insecurity in Campinas, Brazil"}
PT = {"luiz2026", "silva2022a", "cezimbra2022", "werneck2016", "silva2022b", "cezimbra2026b"}
REFN = {}; ORDER = []


def _range(nums):
    out = []; i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
        if j == i: out.append(str(nums[i]))
        elif j == i + 1: out.extend([str(nums[i]), str(nums[j])])
        else: out.append(f"{nums[i]}-{nums[j]}")
        i = j + 1
    return ", ".join(out)


def C(*keys):
    for k in keys:
        if k not in REFN: ORDER.append(k); REFN[k] = len(ORDER)
    return "[" + _range(sorted(REFN[k] for k in keys)) + "]"


def _clean(s): return re.sub(r"<[^>]+>", "", html.unescape(s or "")).strip()


def _ini(g):
    out = ""
    for p in re.split(r"[ \-]", g.strip()):
        if p: out += p[0].upper()
    return out


def _pages(pg):
    if not pg: return ""
    pg = pg.replace("–", "-")
    if "-" in pg:
        a, b = pg.split("-", 1)
        if a.isdigit() and b.isdigit() and len(a) == len(b):
            k = 0
            while k < len(a) - 1 and a[k] == b[k]: k += 1
            bs = b[k:]
            if len(bs) == 1 and len(b) >= 2 and b[-2] == "1": bs = b[-2:]      # 213-15, nao 213-5
            b = bs
        return f"{a}-{b}"
    return pg


KEEP = {"Black", "Brazil", "Brazilian", "Brazilians", "Canada", "American", "Americans", "African", "United", "States", "Rio", "Janeiro", "Latin", "America", "Sweden",
        "Swedish", "Portuguese", "Campinas", "White", "Poisson", "Hunger", "Map", "Salvador", "Belo", "Horizonte", "US", "USA", "UK", "Department", "Agriculture", "Food", "Insecurity"}


def sentence_case(t):
    """so' converte titulos em Title Case; preserva siglas, nomes proprios e o 1o termo"""
    ws = t.split(" ")
    caps = [w for w in ws[1:] if len(w) > 3 and w[0].isupper()]
    if len([w for w in ws[1:] if len(w) > 3]) == 0 or len(caps) / max(1, len([w for w in ws[1:] if len(w) > 3])) < 0.5: return t
    out = [ws[0]]
    for w in ws[1:]:
        core = re.sub(r"^[\W_]+|[\W_]+$", "", w)
        if core in ("Food", "Insecurity") and False: out.append(w); continue
        if (core in KEEP and core not in ("Food", "Insecurity", "Department", "Agriculture", "Hunger", "Map")) or core.isupper() or any(c.isupper() for c in core[1:]) or core.startswith(("U.S", "US")):
            out.append(w)
        else: out.append(w[0].lower() + w[1:] if w and w[0].isalpha() else w.lower())
    return " ".join(out)


def ref_entry(k):
    if k in MANUAL: return MANUAL[k]
    r = META[k]; au = r.get("author", [])
    fam = lambda a: (a.get("family", "").title() if a.get("family", "").isupper() else a.get("family", ""))
    names = [f"{fam(a)} {_ini(a.get('given', ''))}" for a in au[:3]]
    s_ = ", ".join(names) + (" et al" if len(au) > 3 else "")
    title = TITLE_OVR.get(k, _clean(r["title"][0])).rstrip(".")
    title = title if k in TITLE_OVR else sentence_case(title)
    if k in PT: title += " [in Portuguese]"
    jn = _clean(r.get("container-title", [""])[0]); jn = ABREV.get(jn, jn)
    y = ANO.get(k, r["issued"]["date-parts"][0][0])
    vol = r.get("volume"); pg = r.get("page") or r.get("article-number")
    if k == "linkphelan1995": pg = "80-94"
    if k == "heidari2016": pg = "2"
    return f"{s_}. {title}{'' if title.endswith('?') else '.'} *{jn}* {y};{vol or ''}" + (f":{_pages(pg)}" if pg else "") + f". doi:{r['DOI']}"


# ============================================================== Word
WC = {"main": 0, "abstract": 0, "keymsg": 0, "decl": 0, "tables": 0, "captions": 0, "refs": 0, "keywords": 0}
SEC = {}          # palavras do texto principal por secao
CURSEC = {"s": ""}


def wcount(s): return len(re.findall(r"\S+", s))


def _plain(t): return re.sub(r"\*|⟦|⟧", "", t)


def rich(par, text, size=None, bold=None, italic=None):
    """*italico*, **negrito**, ⟦campo a confirmar⟧ (destaque amarelo)"""
    for t in re.split(r"(⟦[^⟧]+⟧|\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not t: continue
        if t.startswith("⟦"): r = par.add_run("[" + t[1:-1] + "]"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif t.startswith("**"): r = par.add_run(t[2:-2]); r.bold = True
        elif t.startswith("*"): r = par.add_run(t[1:-1]); r.italic = True
        else: r = par.add_run(t)
        if size: r.font.size = Pt(size)
        if bold: r.bold = True
        if italic: r.italic = True
    return par


def new_doc(line_numbers=True, spacing=1.5, size=11):
    d = Document()
    cp = d.core_properties; cp.author = ""; cp.last_modified_by = ""; cp.title = ""
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.line_spacing = spacing; st.paragraph_format.space_after = Pt(6)
    sec = d.sections[0]; sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.5)
    if line_numbers:
        ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous")
        sec._sectPr.find(qn("w:pgMar")).addnext(ln)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    def _fld(kind, text=None):
        run = fp.add_run()
        if kind == "instr":
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text
        else:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
        run._r.append(e)
    _fld("begin"); _fld("instr", " PAGE "); _fld("separate"); fp.add_run("1"); _fld("end")
    return d


def H1(d, t, count=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(13)
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.keep_with_next = True
    if count: WC["main"] += wcount(t); SEC[CURSEC["s"]] = SEC.get(CURSEC["s"], 0) + wcount(t)
    return p


def H2(d, t, count=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = True; r.italic = True; r.font.size = Pt(11.5)
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True
    if count: WC["main"] += wcount(t); SEC[CURSEC["s"]] = SEC.get(CURSEC["s"], 0) + wcount(t)
    return p


def para(d, text, bucket="main", align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    text = oxford(text)
    p = d.add_paragraph(); rich(p, text); p.alignment = align
    n = wcount(_plain(text)); WC[bucket] += n
    if bucket == "main": SEC[CURSEC["s"]] = SEC.get(CURSEC["s"], 0) + n
    return p


def caption(d, text, bucket="captions"):
    text = oxford(text)
    p = d.add_paragraph(); rich(p, text, size=9.5); p.paragraph_format.line_spacing = 1.15
    WC[bucket] += wcount(_plain(text)); return p


def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color); tcPr.append(shd)


def _no_vertical_rules(t):
    """IJE: vertical rules prohibited. Apenas linhas horizontais (topo, base, abaixo do cabecalho)."""
    tbl = t._tbl; tblPr = tbl.tblPr
    for e in tblPr.findall(qn("w:tblBorders")): tblPr.remove(e)
    b = OxmlElement("w:tblBorders")
    for edge, val in (("top", "single"), ("left", "nil"), ("bottom", "single"), ("right", "nil"), ("insideH", "nil"), ("insideV", "nil")):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), val)
        if val == "single": el.set(qn("w:sz"), "8"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        b.append(el)
    tblPr.append(b)


def _cell_bottom_border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); tb = OxmlElement("w:tcBorders"); el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000"); tb.append(el); tcPr.append(tb)


def table(d, header, rows, widths=None, font=8, note=None, bucket="tables", group_rows=()):
    t = d.add_table(rows=1, cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_vertical_rules(t)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; p = c.paragraphs[0]; rich(p, oxford(h), size=font, bold=True)
        p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0); _cell_bottom_border(c); WC[bucket] += wcount(_plain(h))
    trPr = t.rows[0]._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells; grp = ri in group_rows
        for i, v in enumerate(row):
            cells[i].text = ""; p = cells[i].paragraphs[0]; p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
            rich(p, oxford(str(v)), size=font, bold=bool(grp))
            if i > 0 and not grp: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            WC[bucket] += wcount(_plain(str(v)))
        if grp:
            for c in cells: set_cell_bg(c, "F2F2F2")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    if note:
        note = oxford(note)
        p = d.add_paragraph(); rich(p, note, size=8); p.paragraph_format.line_spacing = 1.0; WC[bucket] += wcount(_plain(note))
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(d, fname, width_in, cap, bucket="captions", alt=""):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic = p.add_run().add_picture(str(FIG / fname), width=Inches(width_in)); p.paragraph_format.keep_with_next = True
    if alt: pic._inline.docPr.set("descr", alt)
    caption(d, cap, bucket)
