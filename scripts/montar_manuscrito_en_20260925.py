# -*- coding: utf-8 -*-
"""
montar_manuscrito_en_20260925.py

Monta os arquivos do manuscrito em ingles (Social Science & Medicine):
  01_Title_page.docx, 02_Manuscript_blinded.docx, 03_Highlights.docx,
  04_Supplementary_material.docx, 05_Cover_letter.docx, Supplementary_data_strata.xlsx

Todos os numeros vem de ms_numeros_20260925 (CSVs). Todas as referencias vem de
refs_ssm_20260925 (Crossref, DOIs conferidos). Contagem de palavras contra o
limite de 9.000 do SSM (inclui resumo, tabelas, legendas, referencias).
"""
import re, html, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, str(Path(__file__).parent))
from ms_numeros_20260925 import *          # noqa
from refs_ssm_20260925 import DOIS, ANO, SUFIXO, carregar

OUT = BASE / "MANUSCRITO_EN_SSM_20260925"
FIG = OUT / "figures"
META = carregar()

# ============================================================== referencias
MANUAL = {
    "crenshaw1989": ("Crenshaw, 1989", "Crenshaw",
        "Crenshaw, K., 1989. Demarginalizing the intersection of race and sex: a Black feminist critique of antidiscrimination doctrine, feminist theory and antiracist politics. University of Chicago Legal Forum 1989 (1), 139–167."),
    "fao2025": ("FAO et al., 2025", "FAO",
        "FAO, IFAD, UNICEF, WFP, WHO, 2025. The State of Food Security and Nutrition in the World 2025. FAO, Rome. https://doi.org/10.4060/cd6008en"),
    "camara2026": ("Câmara et al., 2026", "Câmara",
        "Câmara, J.H.R., et al., 2026. Racial and territorial inequalities in food insecurity in one of the largest metropolises in Latin America: Rio de Janeiro, Brazil, 2023–2024. American Journal of Public Health, published online ahead of print 23 July 2026. https://doi.org/10.2105/AJPH.2026.308566"),
    "ibge2024": ("IBGE, 2024", "IBGE",
        "IBGE (Instituto Brasileiro de Geografia e Estatística), 2024. Pesquisa Nacional por Amostra de Domicílios Contínua: Segurança Alimentar 2023. IBGE, Rio de Janeiro."),
    "penssan2022": ("Rede PENSSAN, 2022", "Rede PENSSAN",
        "Rede PENSSAN, 2022. II Inquérito Nacional sobre Insegurança Alimentar no Contexto da Pandemia da COVID-19 no Brasil: II VIGISAN. Fundação Friedrich Ebert; Rede PENSSAN, São Paulo."),
}
USED = []


def _clean(s):
    return re.sub(r"<[^>]+>", "", html.unescape(s or "")).strip()


def _fam(a):
    f = a.get("family", "").strip()
    return f.title() if f.isupper() else f


def _year(k):
    return ANO.get(k, META[k]["issued"]["date-parts"][0][0])


def label(k):
    if k in MANUAL: return MANUAL[k][0]
    au = META[k].get("author", [])
    y = f"{_year(k)}{SUFIXO.get(k, '')}"
    if len(au) == 1: return f"{_fam(au[0])}, {y}"
    if len(au) == 2: return f"{_fam(au[0])} and {_fam(au[1])}, {y}"
    return f"{_fam(au[0])} et al., {y}"


def narr(k):
    """citacao narrativa: 'Luiz et al. (2026)'"""
    l = label(k)
    a, y = l.rsplit(", ", 1)
    if k not in USED: USED.append(k)
    return f"{a} ({y})"


def P(*keys):
    for k in keys:
        if k not in USED: USED.append(k)
    return "(" + "; ".join(label(k) for k in keys) + ")"


def _ini(given):
    parts = re.split(r"([ \-])", given.strip())
    out = ""
    for p in parts:
        if p == "-": out += "-"
        elif p == " " or not p: continue
        else: out += p[0].upper() + "."
    return out


def ref_entry(k):
    if k in MANUAL: return MANUAL[k][2]
    r = META[k]
    au = r.get("author", [])
    names = [f"{_fam(a)}, {_ini(a.get('given', ''))}" for a in au[:6]]
    s = ", ".join(names) + (", et al." if len(au) > 6 else "")
    y = _year(k)
    if k in SUFIXO: y = f"{y}{SUFIXO[k]}"
    title = _clean(r["title"][0]).rstrip(".")
    jn = _clean(r.get("container-title", [""])[0])
    vol, iss = r.get("volume"), r.get("issue")
    pg = r.get("page") or r.get("article-number")
    if k == "linkphelan1995": pg = "80–94"
    if k == "heidari2016": pg = "2"
    va = f"{vol}" + (f" ({iss})" if iss else "")
    pg = pg.replace("-", "–") if pg else ""
    return f"{s}, {y}. {title}. {jn} {va}" + (f", {pg}" if pg else "") + f". https://doi.org/{r['DOI']}"


# ============================================================== docx utils
WC = {"abstract": 0, "main": 0, "tables": 0, "captions": 0, "refs": 0, "decl": 0}


def wcount(s): return len(re.findall(r"\S+", s))


def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)


def new_doc(line_numbers=False, spacing=1.5, font="Times New Roman", size=11):
    d = Document()
    cp = d.core_properties                     # anonimato: sem autor/titulo nos metadados
    cp.author = ""; cp.last_modified_by = ""; cp.title = ""; cp.comments = ""; cp.keywords = ""
    st = d.styles["Normal"]
    st.font.name = font; st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.paragraph_format.line_spacing = spacing
    st.paragraph_format.space_after = Pt(6)
    sec = d.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.5)
    if line_numbers:
        ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous")
        pgmar = sec._sectPr.find(qn("w:pgMar"))          # ordem do esquema: lnNumType vem logo apos pgMar
        pgmar.addnext(ln)
    # numero de pagina
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    def _fld(kind, text=None):
        run = fp.add_run()
        if kind == "instr":
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text
        else:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
        run._r.append(e)
    _fld("begin"); _fld("instr", " PAGE "); _fld("separate")
    fp.add_run("1")
    _fld("end")
    return d


def rich(par, text, size=None, bold=None, italic=None):
    """*italico* e **negrito** minimos."""
    toks = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for t in toks:
        if not t: continue
        if t.startswith("**"): r = par.add_run(t[2:-2]); r.bold = True
        elif t.startswith("*"): r = par.add_run(t[1:-1]); r.italic = True
        else: r = par.add_run(t)
        if size: r.font.size = Pt(size)
        if bold: r.bold = True
        if italic: r.italic = True
    return par


def H1(d, t, bucket=None):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(13)
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.keep_with_next = True
    return p


def H2(d, t):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = True; r.italic = True; r.font.size = Pt(11.5)
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True
    return p


def para(d, text, bucket="main", align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=False):
    p = d.add_paragraph(); rich(p, text); p.alignment = align
    if first_indent: p.paragraph_format.first_line_indent = Cm(0.9)
    WC[bucket] += wcount(re.sub(r"\*", "", text))
    return p


def caption(d, text, bucket="captions"):
    p = d.add_paragraph(); rich(p, text, size=9.5); p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = False
    WC[bucket] += wcount(re.sub(r"\*", "", text))
    return p


def table(d, header, rows, widths=None, font=8, bold_first_col=False, note=None, bucket="tables", group_rows=()):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; p = c.paragraphs[0]; r = p.add_run(h); r.bold = True; r.font.size = Pt(font)
        p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
        set_cell_bg(c, "E7E6E6")
        WC[bucket] += wcount(h)
    set_repeat_header(t.rows[0])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        grp = ri in group_rows
        for i, v in enumerate(row):
            cells[i].text = ""; p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(v)); r.font.size = Pt(font)
            if grp or (bold_first_col and i == 0): r.bold = grp
            if i > 0 and not grp: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            WC[bucket] += wcount(str(v))
        if grp:
            for c in cells: set_cell_bg(c, "F2F2F2")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    if note:
        p = d.add_paragraph(); rich(p, note, size=8); p.paragraph_format.line_spacing = 1.0
        WC[bucket] += wcount(re.sub(r"\*", "", note))
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(d, fname, width_in, cap):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG / fname), width=Inches(width_in))
    p.paragraph_format.keep_with_next = True
    caption(d, cap)


# ============================================================== numeros do texto
tot = T1("Total", "Brasil")
RAW = 473206                       # registros de pessoas lidos (log da analise)
NA = int(tot.n)
prev_any = f"{f1(tot.ia_total_pct)}% (95% CI {ci(tot.ia_total_lo, tot.ia_total_hi, 1)})"
prev_sev = f"{f1(tot.ia_grave_pct)}% (95% CI {ci(tot.ia_grave_lo, tot.ia_grave_hi, 1)})"
black = T1("Cor/raca", "negra (preta+parda)"); nonblack = T1("Cor/raca", "nao negra")
woman = T1("Sexo", "mulher"); rural = T1("Situacao", "rural")
edu1 = T1("Instrucao", "sem_fund"); inc1 = T1("Renda pc", "q1"); inc5 = T1("Renda pc", "q5")
edu4 = T1("Instrucao", "superior")

a2 = {k: ATEN(*k) for k in [("Negra", "bruto", "ia_total"), ("Negra", "renda_e_demais", "ia_total")]}
def AT(var, esp, d): return ATEN(var, esp, d)

# pares
PS, PE, PI, PR = "Sex", "Education", "Income", "Residence"
pair_key = {"sex": "Race x Sex", "edu": "Race x Education", "inc": "Race x Income", "res": "Race x Residence"}
def PAIR(k, d): return T2(pair_key[k], d)

# MAIHDA
M = {d: m1.loc[d] for d in ["ia_total", "ia_grave"]}
def coef(d, termo):
    return m2[(m2.desfecho == d) & (m2.termo == termo)].iloc[0]
def strata_ext(d):
    s = m3[m3.desfecho == d].sort_values("p_completo")
    return s.iloc[0], s.iloc[-1]
lo_any, hi_any = strata_ext("ia_total"); lo_sev, hi_sev = strata_ext("ia_grave")

# ============================================================== TEXTO
def W(n):
    """numero por extenso ate dez (estilo de periodico)"""
    return {0: "zero", 1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six", 7: "seven", 8: "eight", 9: "nine", 10: "ten"}.get(int(n), str(int(n)))


def jd(par, desf):
    """decomposicao da disparidade conjunta na escala de risco (Jackson et al., 2016), em pontos percentuais"""
    r = PAIR(par, desf)
    joint = r.p11 - r.p00; a = r.p10 - r.p00; b = r.p01 - r.p00; inter = r.p11 - r.p10 - r.p01 + r.p00
    return joint, a, b, inter


def est_cib(x, lo, hi, dec=2):
    """estimativa com IC entre colchetes, para uso dentro de parenteses"""
    fmt = {1: f1, 2: f2, 3: f3}[dec]
    return f"{fmt(x)} [95% CI {ci(lo, hi, dec)}]"


def build_abstract():
    p11s = PAIR("edu", "ia_grave"); pi_s = PAIR("inc", "ia_grave")
    r_lo = T6("negra", "renda", "q1", "ia_total"); r_hi = T6("negra", "renda", "q5", "ia_total")
    ab = (
        f"Race, gender, education, income and rurality each pattern food insecurity in Brazil, but whether their combination exceeds the sum "
        f"or the product of its parts depends on the scale of analysis. This cross-sectional study analysed {n0(NA)} households from the Continuous "
        f"National Household Sample Survey (fourth quarter 2023) with the Brazilian Food Insecurity Scale, characterising households by their reference "
        f"person. Methods were Poisson models with robust variance, interaction on additive (relative excess risk due to interaction, RERI) and "
        f"multiplicative (ratio of risk ratios) scales, cumulative-disadvantage gradients, and intersectional multilevel analysis of individual "
        f"heterogeneity and discriminatory accuracy (MAIHDA) across 160 strata. Prevalence was {f1(tot.ia_total_pct)}% (any) and {f1(tot.ia_grave_pct)}% (severe). "
        f"Income accounted for about half of the crude associations of race, sex and education; Black race remained associated with any food insecurity "
        f"(adjusted risk ratio {f2(AT('Negra','renda_e_demais','ia_total').rr)}). For severe food insecurity, race-by-education and race-by-income interactions "
        f"were super-additive (RERI {f2(p11s.reri)} and {f2(pi_s.reri)}) yet sub-multiplicative (ratios {f2(p11s.ror)} and {f2(pi_s.ror)}). The relative racial gap "
        f"was larger at higher income (risk ratio {f2(r_lo.rr_ajustado)} vs {f2(r_hi.rr_ajustado)}) but the absolute gap smaller ({f1(r_lo.dif_pp)} vs {f1(r_hi.dif_pp)} "
        f"percentage points). Strata carried {f1(M['ia_total'].vpc_nulo*100)}% and {f1(M['ia_grave'].vpc_nulo*100)}% of latent variance, of which additive main effects "
        f"explained {f1(M['ia_total'].pcv*100)}% and {f1(M['ia_grave'].pcv*100)}%; discriminatory accuracy was moderate (AUC {f2(AUC['ia_total']['p_completo'])}, "
        f"{f2(AUC['ia_grave']['p_completo'])}). Intersectional inequality in food insecurity in Brazil is large yet largely the accumulation of independent "
        f"disadvantages, dominated by income; racial disadvantage persists at every income level. Reporting both interaction scales avoids over-reading "
        f"excess risk as synergy."
    )
    return ab


def build_intro(d):
    para(d, f"Food insecurity, the limited or uncertain access to adequate food, is consistently associated with worse physical and mental health across the life course {P('gundersen2015')}. Its recent history in Brazil has been volatile. Food security fell between 2013 and 2018 {P('santos2023')}, hunger rose during the COVID-19 pandemic {P('penssan2022')}, and undernourishment then declined enough for the country to leave the FAO Hunger Map {P('fao2025')}. Even so, 21.6 million households, 27.6% of the total, reported some degree of food insecurity in 2023 {P('ibge2024')}. Progress in the aggregate says little about who remains behind.")
    para(d, f"Food insecurity is a useful case for intersectional analysis. It is measured at the household level, so social positions must be assigned through a reference person and the link to individual gender is indirect {P('broussard2019')}. Its severe form is relatively uncommon ({f1(tot.ia_grave_pct)}% of households here), so ratios and differences can tell different stories. And its gradient by income is steep, which makes the choice of scale consequential: when one exposure carries a very large risk ratio, joint excess risk can arise mechanically.")
    para(d, f"In Brazil, as elsewhere, food insecurity is unequally distributed by income, education, race, sex and place of residence. Household surveys repeatedly find higher prevalence where the reference person is Black or Brown, a woman or poor {P('santos2023','miguel2025')}. Racism is increasingly recognised in Brazil as a social determinant of health, operating through structural, institutional and symbolic processes {P('werneck2016','costa2026')}. Fundamental-cause theory holds that social conditions such as income, education and racism affect health through many changing mechanisms and therefore persist as risk factors even as specific pathways are addressed {P('linkphelan1995','phelanlink2015')}. Empirically, socioeconomic position accounts for a large but incomplete part of racial health gaps, and the returns to education and income are often smaller for racialised groups, a pattern known as diminished returns {P('farmer2005','assari2018','bakhtiari2022')}. For sex and gender, global evidence shows women more likely than men to be food insecure, with income, education and social networks explaining most of the gap in many settings {P('broussard2019','smith2017','ivers2011')} but not in all: in Chile it tracked marital status and household composition rather than income or education {P('silva2023')}.")
    para(d, f"Intersectionality, formulated by Black feminist scholars and named by {narr('crenshaw1989')} and {narr('crenshaw1991')}, holds that categories such as race and gender do not act independently: the position of a Black woman cannot be understood by adding what is known about Black people and about women. Its uptake in population health has been substantial {P('bowleg2012')}, but its translation into quantitative research is contested. A systematic review of 707 quantitative studies found that core theoretical tenets are often lost or misinterpreted and that methods were often simple {P('bauer2021')}. {narr('bauer2014')} warned in particular against conflating metaphorical language with statistical interaction and about the choice of scale.")
    para(d, f"Whether an intersectional disparity is more than the sum of its parts depends on what is summed. Interaction on the additive scale of risk, the excess of joint risk over the sum of the separate excesses, bears directly on public-health impact, whereas interaction on the multiplicative scale is what regression coefficients on log or logit scales test. The two can disagree, and reporting both is recommended {P('rothman1980','knol2012','vanderweele2014')}. {narr('jackson2016')} add that the joint disparity between the most and least advantaged groups can be decomposed into contributions of each axis and their interaction. A different route is intersectional multilevel analysis of individual heterogeneity and discriminatory accuracy (MAIHDA), which treats intersectional strata as a random effect, quantifies how much individual outcomes vary between strata, and tests whether additive main effects suffice to describe that variation {P('evans2018','merlo2018','evans2024a')}. In simulations MAIHDA was the most accurate method for estimating stratum-specific prevalence in smaller samples, although without advantage over main-effects regression {P('mahendran2022')}. The approach is not uncontested: critics argue that strata are individual-level composites, so that the demographic variables enter at both levels, and favour single-level models with interaction terms {P('wilkes2024')}, a view its proponents rebut {P('evans2024b')}; whether stratum residuals can be read as interaction effects is also debated {P('lizotte2020','evans2020')}, and shrinkage implicitly reweights strata towards equal size {P('bashir2026')}. A scoping review of quantitative nutrition research found 55 studies of intersecting inequalities, mostly on body weight and from the United States; only two of 15 studies using nonlinear models assessed interaction on the additive scale {P('fivian2024')}.")
    para(d, f"Brazilian studies have begun to examine intersections in food insecurity. Analyses of national budget and household surveys profile households by the sex and race of the reference person and find the highest prevalence among Black and Brown women {P('santos2022','santos2023','miguel2025')}, with the highest risk in households headed by Black or Brown single mothers of young children {P('santos2023')}. Similar findings come from a pandemic survey {P('luiz2025')}, from Salvador {P('silva2022a')}, from Rio de Janeiro, where racial disparities persisted across income strata {P('camara2026')}, and from Belo Horizonte, where, among people in unfavourable socioeconomic conditions, Black women had the highest odds of food insecurity in interaction models (odds ratio 7.5) {P('ferreira2025')}. {narr('luiz2026')} traced six sex-by-race profiles from 2013 to 2023, using the same 2023 survey analysed here. Intersectional analyses of other outcomes in Brazil report that combinations of race, education and wealth carry pronounced effects for AIDS among 28.3 million people {P('lua2026')}, and, in a Brazilian MAIHDA of obesity, that main effects explained only 46% of between-stratum variance {P('fanton2026')}. Three gaps remain. First, these studies treat sex and race as the intersecting pair, and income, education and residence as covariates or strata rather than as axes that themselves intersect with race. Second, none partitions the disparity into additive and interactive components or reports interaction on both scales. Third, the validity of comparing such groups on a single scale has been questioned {P('silva2022b','cezimbra2022')}, although evidence from the 2013 national survey supports approximate measurement invariance across sex, race and education {P('cezimbra2026a')}.")
    para(d, "Two readings of intersectional inequality are therefore possible. If disadvantage were synergistic, race-by-income and race-by-education interactions would be positive on both scales and between-strata variance would not be explained by main effects. If it were mainly cumulative, interactions on the multiplicative scale would be null or negative, additive-scale excess would be modest relative to what multiplicativity implies, and main effects would account for nearly all between-strata variance.")
    para(d, "We used the 2023 national household survey to ask three questions about household food insecurity in Brazil. First, do race, sex, education and rural residence retain associations after accounting for household income? Second, does race interact with each of the other four dimensions, on the additive and multiplicative scales, and do these interactions survive mutual adjustment? Third, across the 160 intersectional strata defined by all five dimensions, how much between-strata variation is explained by additive main effects, and how well do strata discriminate individual households?")


def build_methods(d):
    H2(d, "2.1 Data and study population")
    para(d, f"We analysed public-use microdata from the fourth quarter of 2023 of the Continuous National Household Sample Survey (Pesquisa Nacional por Amostra de Domicílios Contínua, PNADC), a probability survey of Brazilian households conducted by the Brazilian Institute of Geography and Statistics (IBGE) with a stratified, clustered design, which in that quarter included a food security supplement {P('ibge2024')}. The unit of analysis is the household, described by the characteristics of its reference person, the member recognised by the others as responsible for it. Of {n0(RAW)} person records read, {n0(NA)} households had a valid food insecurity classification and complete data on all variables and formed the analytic sample; no household was excluded because of missing covariates. We followed the STROBE guidelines {P('vonelm2007')} (Supplementary Table S1).")
    H2(d, "2.2 Outcomes")
    para(d, f"IBGE classified households with the Brazilian Food Insecurity Scale (Escala Brasileira de Insegurança Alimentar, EBIA), an experiential household-level scale validated in Brazil {P('perezescamilla2004')}, as secure or mildly, moderately or severely insecure. We analysed any food insecurity (mild, moderate or severe) and severe food insecurity (the most severe category), using the classification distributed with the microdata.")
    H2(d, "2.3 Social positions")
    para(d, f"We examined five characteristics of the reference person or household. (1) Race/colour, self-reported in IBGE categories and grouped as Black (pretos and pardos, that is, Black and Brown people; hereafter Black, following IBGE and Brazilian health-equity literature) versus all others (white, Asian and Indigenous), because small numbers precluded separate analysis of the latter. (2) Sex, recorded in the survey as woman or man; we treat it as a marker of gender relations in household headship in line with SAGER guidance {P('heidari2016')}, while acknowledging that gender identity is not measured. (3) Education in four categories: no schooling or incomplete primary, complete primary or incomplete secondary, complete secondary, and tertiary. (4) Per capita household income in five bands of the monthly minimum wage (MW): ≤1/4, >1/4–1/2, >1/2–1, >1–2 and >2. (5) Residence, urban or rural.")
    H2(d, "2.4 Statistical analysis")
    para(d, f"Except for the multilevel models, we used Poisson regression with robust variance to estimate risk ratios {P('zou2004','zou2013')}, with survey weights rescaled to mean one and standard errors clustered on primary sampling units. We preferred Poisson to logistic regression because the outcome is common, so odds ratios would overstate risk ratios. Weighted prevalences carry 95% confidence intervals from Taylor linearisation at the primary-sampling-unit level. The public microdata lack the stratification variable, so variance estimates ignore stratification, which is generally conservative.")
    para(d, "*Own associations.* For race, sex, education (no schooling or incomplete primary versus all others) and rural residence we estimated the crude risk ratio, the ratio adjusted for income band, and the ratio adjusted for income and the other three characteristics, and the share of the crude log risk ratio attenuated by adjustment.")
    para(d, f"*Pairwise interaction.* For race with each other characteristic, dichotomised against all others (woman; no schooling or incomplete primary; income ≤1/4 MW; rural), we estimated risk ratios for the three exposed joint categories against the doubly unexposed reference, and reported interaction on both scales {P('knol2012')}. On the additive scale we used the relative excess risk due to interaction, RERI = RR11 − RR10 − RR01 + 1 {P('rothman1980','vanderweele2014')}, with 95% confidence intervals by the delta method using the cluster-robust covariance {P('andersson2005')}, confirmed by a cluster bootstrap of 300 replicates (Supplementary Table S2). On the multiplicative scale we report the ratio of risk ratios, RR11/(RR10 × RR01), the exponentiated product-term coefficient. Because RERI is positive whenever two risk ratios exceed one, even with purely multiplicative joint effects, we also report the RERI expected under multiplicativity, (RR10 − 1)(RR01 − 1), as a benchmark.")
    para(d, "*Mutual adjustment and strata.* One model included the five main effects and the four race-by-characteristic products, with a joint Wald test (4 degrees of freedom) of the products. Following recommendations to present effects within strata of the other factor, we estimated the adjusted risk ratio and the risk difference for Black race within income bands and education categories, and for sex within income bands and race groups, adjusting for the remaining characteristics. Three-way interactions (race × sex × income; race × sex × education) were tested in separate models.")
    para(d, "*Cumulative disadvantage.* We counted, for each household, how many of four disadvantages it had (woman reference person; no schooling or incomplete primary; income ≤1/4 MW; rural residence; range 0–4). By race we estimated the prevalence at each count and the change per additional disadvantage on the relative (Poisson) and absolute (linear probability model) scales, testing whether slopes differed by race.")
    para(d, f"*Intersectional MAIHDA.* The cross-classification of race (2), sex (2), education (4), income (5) and residence (2) defines 160 strata, all occupied (median {n0(N_STRATA.median())} households, range {n0(N_STRATA.min())}–{n0(N_STRATA.max())}; {int((N_STRATA < 50).sum())} strata with fewer than 50 households). Following {narr('evans2024a')} and {narr('merlo2018')}, we fitted logistic models with a stratum random intercept: Model A (null) and Model B, adding the five characteristics as fixed effects. The variance partition coefficient (VPC) is the share of latent variance lying between strata, σ²/(σ² + π²/3), with a profile-likelihood interval in the null model. The proportional change in stratum variance (PCV) from A to B, (σ²A − σ²B)/σ²A, is the share of between-stratum variance explained by additive main effects; the remainder reflects interaction. Discriminatory accuracy was the area under the receiver operating characteristic curve (AUC) of stratum-predicted probabilities. To find strata departing from additivity we examined Model B stratum residuals with their conditional standard errors, expecting about 5% to exceed |z| = 1.96 by chance, and computed, on the risk scale, the difference between full-model and main-effects-only predicted prevalence. Because MAIHDA is debated {P('wilkes2024','evans2024b','bashir2026')}, we treat it as descriptive and complementary to the single-level interaction models, and read stratum residuals as departures from additivity {P('evans2020')}, an interpretation that has been contested {P('lizotte2020')}. As in the standard procedure, the models do not use survey weights; as sensitivity analysis we refitted them with normalised weights, an approximation that is not design-consistent.")
    para(d, "*Status of the analyses.* The plan developed iteratively. The own-association, pairwise and joint analyses preceded the multilevel analysis, whereas the benchmark for multiplicative RERI, the stratum-specific gaps and the absolute-scale gradients were added after inspecting initial results and should be regarded as exploratory. The study was not preregistered.")
    H2(d, "2.5 Software and ethics")
    para(d, "Analyses used Python 3.12 (statsmodels 0.14.6) and R 4.6.1 (lme4 2.0.6). The study used publicly available, anonymised secondary data and, under Brazilian regulation (Resolution CNS 510/2016), did not require ethics committee review.")


def build_results(d):
    H2(d, "3.1 Sample and prevalence")
    para(d, f"The analytic sample comprised {n0(NA)} households. Weighted prevalence was {prev_any} for any and {prev_sev} for severe food insecurity. {f1(black.pct_pond_amostra)}% of reference persons were Black, {f1(woman.pct_pond_amostra)}% women and {f1(rural.pct_pond_amostra)}% rural residents; {f1(edu1.pct_pond_amostra)}% had no schooling or incomplete primary education and {f1(inc1.pct_pond_amostra)}% had income at or below 1/4 MW per capita (Table 1). Any food insecurity ranged from {f1(inc1.ia_total_pct)}% in the lowest to {f1(inc5.ia_total_pct)}% in the highest income band, and from {f1(edu1.ia_total_pct)}% to {f1(edu4.ia_total_pct)}% across education; it was {f1(black.ia_total_pct)}% in Black and {f1(nonblack.ia_total_pct)}% in non-Black households.")
    table_1(d)
    H2(d, "3.2 Own associations and the role of income")
    b, bi, bf = AT("Negra", "bruto", "ia_total"), AT("Negra", "so_renda", "ia_total"), AT("Negra", "renda_e_demais", "ia_total")
    w_, wi = AT("Mulher", "bruto", "ia_total"), AT("Mulher", "renda_e_demais", "ia_total")
    ws_, wsi = AT("Mulher", "so_renda", "ia_grave"), AT("Mulher", "renda_e_demais", "ia_grave")
    ru, rui, ruf = AT("Rural", "bruto", "ia_total"), AT("Rural", "so_renda", "ia_total"), AT("Rural", "renda_e_demais", "ia_total")
    rs, rsf = AT("Rural", "bruto", "ia_grave"), AT("Rural", "renda_e_demais", "ia_grave")
    para(d, f"Each characteristic was associated with food insecurity in crude analysis (Table 2). Adjusting for income halved the crude log risk ratio for Black race ({f2(b.rr)} to {f2(bi.rr)}; {f1(bi.atenuacao_pct)}% attenuation), and the association remained after further adjustment ({est_cib(bf.rr, bf.ic95_low, bf.ic95_high)}). The pattern was similar for women, whose ratio for any food insecurity fell from {f2(w_.rr)} to {f2(wi.rr)} after full adjustment, and for education. For severe food insecurity, income adjustment alone reduced the sex association to {est_ci(ws_.rr, ws_.ic95_low, ws_.ic95_high)}, and full adjustment left {est_ci(wsi.rr, wsi.ic95_low, wsi.ic95_high)}. Rural residence, a risk factor in crude analysis (risk ratio {f2(ru.rr)} for any and {f2(rs.rr)} for severe food insecurity), reversed after adjustment for income ({f2(rui.rr)} for any food insecurity) and further with full adjustment ({f2(ruf.rr)} for any and {f2(rsf.rr)} for severe), because rural households are poorer.")
    table_2(d)
    H2(d, "3.3 Pairwise interactions on two scales")
    s_, e_, i_, r_ = PAIR("sex", "ia_grave"), PAIR("edu", "ia_grave"), PAIR("inc", "ia_grave"), PAIR("res", "ia_grave")
    sa, ea, ia, ra = PAIR("sex", "ia_total"), PAIR("edu", "ia_total"), PAIR("inc", "ia_total"), PAIR("res", "ia_total")
    n_rer_sig = int(((t2.reri_lo_delta > 0)).sum())
    n_ror_lt1 = int((t2.ror < 1).sum()); n_ror_sig = int((t2.ror_hi < 1).sum())
    n_below = int((t2.reri < t2.reri_esperado_nulo_mult).sum())
    para(d, f"Joint exposure carried the highest risks. Black households with income at or below 1/4 MW had {f1(i_.p11)}% severe food insecurity, against {f1(i_.p00)}% in non-Black households above that income, a risk ratio of {est_ci(i_.rr11, i_.rr11_lo, i_.rr11_hi)} (Table 3, Fig. 1). RERI was positive in all eight comparisons and its confidence interval excluded zero in {W(n_rer_sig)}: for sex (any, {est_cib(sa.reri, sa.reri_lo_delta, sa.reri_hi_delta)}; severe, {est_cib(s_.reri, s_.reri_lo_delta, s_.reri_hi_delta)}), for education and income in severe food insecurity ({est_cib(e_.reri, e_.reri_lo_delta, e_.reri_hi_delta)}; {est_cib(i_.reri, i_.reri_lo_delta, i_.reri_hi_delta)}), and for residence in any food insecurity ({est_cib(ra.reri, ra.reri_lo_delta, ra.reri_hi_delta)}). For any food insecurity the intervals for education ({ci(ea.reri_lo_delta, ea.reri_hi_delta)}) and income ({ci(ia.reri_lo_delta, ia.reri_hi_delta)}) included zero.")
    para(d, f"The multiplicative scale told a different story. The ratio of risk ratios was below one in {W(n_ror_lt1)} of eight comparisons and significantly so in {W(n_ror_sig)}: education ({est_cib(ea.ror, ea.ror_lo, ea.ror_hi)} for any and {est_cib(e_.ror, e_.ror_lo, e_.ror_hi)} for severe food insecurity), income ({est_cib(ia.ror, ia.ror_lo, ia.ror_hi)} and {est_cib(i_.ror, i_.ror_lo, i_.ror_hi)}) and, for any food insecurity, sex ({est_cib(sa.ror, sa.ror_lo, sa.ror_hi)}). For residence it was close to one ({est_cib(ra.ror, ra.ror_lo, ra.ror_hi)} for any food insecurity). In {W(n_below)} of eight comparisons the observed RERI was smaller than that expected under multiplicativity; for race and income in severe food insecurity it was {f2(i_.reri)} against an expected {f2(i_.reri_esperado_nulo_mult)}. Joint effects were therefore super-additive but sub-multiplicative.")
    jd_i, jd_e = {dd: jd("inc", dd) for dd in ["ia_total", "ia_grave"]}, {dd: jd("edu", dd) for dd in ["ia_total", "ia_grave"]}
    def _jd(t, other): return f"{f1(t[0])} percentage points: {f1(t[1])} attributable to race alone, {f1(t[2])} to {other} alone and {f1(t[3])} ({f1(t[3]/t[0]*100)}%) to their interaction"
    para(d, f"On the risk scale (Supplementary Table S2), the joint disparity between Black households with income at or below 1/4 MW and non-Black households above it was, for any food insecurity, {_jd(jd_i['ia_total'], 'income')}. For severe food insecurity it was {_jd(jd_i['ia_grave'], 'income')}, the excess intersectional disparity {P('jackson2016')}. For race and education, the joint disparity in severe food insecurity was {_jd(jd_e['ia_grave'], 'education')}.")
    rg = reg[reg.parte == "b_regional"]
    nm = {"Norte": "North", "Nordeste": "Northeast", "Centro-Oeste": "Central-West", "Sudeste": "Southeast", "Sul": "South"}
    pos = rg[rg.reri > 0]; sig = rg[rg.reri_ic95_low > 0]
    sig_txt = " and ".join(f"the {nm[r.regiao]} ({'any' if r.desfecho == 'ia_total' else 'severe'} food insecurity)" for r in sig.itertuples())
    para(d, f"For race × sex by macro-region (Supplementary Table S3), RERI was positive in {W(len(pos))} of {W(len(rg))} region-outcome combinations, and its bootstrap interval excluded zero in {W(len(sig))}: {sig_txt}.")
    table_3(d)
    figure(d, "fig1_interaction_scales.png", 6.3, "**Fig. 1.** Interaction between race and four characteristics on additive and multiplicative scales. (a) RERI with 95% confidence interval (delta method); diamonds show the RERI expected if joint effects were purely multiplicative. (b) Ratio of risk ratios with 95% confidence interval. Each characteristic is dichotomised against all others: woman; no schooling or incomplete primary education; income ≤1/4 minimum wage per capita; rural residence.")
    H2(d, "3.4 Mutually adjusted model and three-way interactions")
    j = {(t, dd): T3(t, dd) for t in ["negra:mulher", "negra:sem_fund", "negra:renda_q1", "negra:rural"] for dd in ["ia_total", "ia_grave"]}
    wa, ws = T3("WALD_conjunto_4gl", "ia_total"), T3("WALD_conjunto_4gl", "ia_grave")
    para(d, f"With the four products in one model, the multiplicative interactions for education and income remained below one for any ({est_cib(j[('negra:sem_fund','ia_total')].rr, j[('negra:sem_fund','ia_total')].lo, j[('negra:sem_fund','ia_total')].hi)} and {est_cib(j[('negra:renda_q1','ia_total')].rr, j[('negra:renda_q1','ia_total')].lo, j[('negra:renda_q1','ia_total')].hi)}) and severe food insecurity ({est_cib(j[('negra:sem_fund','ia_grave')].rr, j[('negra:sem_fund','ia_grave')].lo, j[('negra:sem_fund','ia_grave')].hi)} and {est_cib(j[('negra:renda_q1','ia_grave')].rr, j[('negra:renda_q1','ia_grave')].lo, j[('negra:renda_q1','ia_grave')].hi)}), whereas sex and residence products were near one (Table 4). The joint test of the four products was significant for both outcomes (Wald χ² {f1(wa.rr)} and {f1(ws.rr)}, 4 df, both p<0.001). Three-way interactions were not detected: for race × sex × income the ratio of risk ratios was {est_ci(T7('renda_q1','ia_total').rr, T7('renda_q1','ia_total').lo, T7('renda_q1','ia_total').hi)} for any and {est_ci(T7('renda_q1','ia_grave').rr, T7('renda_q1','ia_grave').lo, T7('renda_q1','ia_grave').hi)} for severe food insecurity; for race × sex × education, {est_ci(T7('sem_fund','ia_total').rr, T7('sem_fund','ia_total').lo, T7('sem_fund','ia_total').hi)} and {est_ci(T7('sem_fund','ia_grave').rr, T7('sem_fund','ia_grave').lo, T7('sem_fund','ia_grave').hi)} (Supplementary Table S4).")
    table_4(d)
    H2(d, "3.5 Race and sex gaps within strata")
    q1, q5 = T6("negra", "renda", "q1", "ia_total"), T6("negra", "renda", "q5", "ia_total")
    s1, s5 = T6("negra", "renda", "q1", "ia_grave"), T6("negra", "renda", "q5", "ia_grave")
    e1, e4 = T6("negra", "instrucao", "sem_fund", "ia_total"), T6("negra", "instrucao", "superior", "ia_total")
    sx = [T6("mulher", "renda", q, "ia_total") for q in ["q1", "q2", "q3", "q4", "q5"]]
    sxs = [T6("mulher", "renda", q, "ia_grave") for q in ["q1", "q2", "q3", "q4", "q5"]]
    xn, xb = T6("mulher", "raca", "nao_negra", "ia_total"), T6("mulher", "raca", "negra", "ia_total")
    para(d, f"Within income bands, the adjusted risk ratio for Black race was {est_ci(q1.rr_ajustado, q1.lo, q1.hi)} in the lowest band and {est_ci(q5.rr_ajustado, q5.lo, q5.hi)} in the highest for any food insecurity, and {est_ci(s1.rr_ajustado, s1.lo, s1.hi)} and {est_ci(s5.rr_ajustado, s5.lo, s5.hi)} for severe food insecurity (Fig. 2). The absolute gap moved the other way, from {f1(q1.dif_pp)} percentage points (95% CI {ci(q1.dif_pp_lo, q1.dif_pp_hi, 1)}) in the lowest band to {f1(q5.dif_pp)} ({ci(q5.dif_pp_lo, q5.dif_pp_hi, 1)}) in the highest. By education the risk ratio for any food insecurity was {f2(e1.rr_ajustado)} among those with no schooling or incomplete primary education and {f2(e4.rr_ajustado)} among those with tertiary education, with absolute gaps of {f1(e1.dif_pp)} and {f1(e4.dif_pp)} points (Supplementary Fig. S1). The adjusted risk ratio for women and any food insecurity ranged from {f2(min(x.rr_ajustado for x in sx))} to {f2(max(x.rr_ajustado for x in sx))} across income bands; for severe food insecurity every band interval included one (Supplementary Table S5). By race, it was {est_ci(xn.rr_ajustado, xn.lo, xn.hi)} among non-Black and {est_ci(xb.rr_ajustado, xb.lo, xb.hi)} among Black households.")
    figure(d, "fig2_race_gaps_income.png", 6.0, "**Fig. 2.** Racial gap in food insecurity across income bands. Adjusted risk ratio (a, c; adjusted for sex, education and residence) and absolute difference in prevalence (b, d; crude) for Black versus non-Black reference persons within bands of per capita household income (minimum wages), for any (a, b) and severe (c, d) food insecurity. Bars are 95% confidence intervals.")
    H2(d, "3.6 Cumulative disadvantage")
    g0a, g4a = T4("ia_total", 0), T4("ia_total", 4)
    g0s, g4s = T4("ia_grave", 0), T4("ia_grave", 4)
    r_rel, r_abs = T5("ia_total", "relativa"), T5("ia_total", "absoluta")
    rs_rel, rs_abs = T5("ia_grave", "relativa"), T5("ia_grave", "absoluta")
    para(d, f"Prevalence of any food insecurity rose from {f1(g0a.nao_negra_pct)}% with no disadvantage to {f1(g4a.nao_negra_pct)}% with all four in non-Black households, and from {f1(g0a.negra_pct)}% to {f1(g4a.negra_pct)}% in Black households (Fig. 3); for severe food insecurity, from {f1(g0s.nao_negra_pct)}% to {f1(g4s.nao_negra_pct)}% and from {f1(g0s.negra_pct)}% to {f1(g4s.negra_pct)}%. Each additional disadvantage multiplied the risk of any food insecurity by {f2(r_rel.nao_negra)} in non-Black and {f2(r_rel.negra)} in Black households (difference p<0.001), a flatter relative slope among Black households, but added {f1(r_abs.nao_negra)} percentage points in non-Black and {f1(r_abs.negra)} in Black households (p<0.001), a steeper absolute slope. Results for severe food insecurity were similar (relative {f2(rs_rel.nao_negra)} versus {f2(rs_rel.negra)}; absolute {f1(rs_abs.nao_negra)} versus {f1(rs_abs.negra)} points). Only {n0(g4a.nao_negra_n)} non-Black households had all four disadvantages, so estimates at the top of the gradient are imprecise (Supplementary Table S6).")
    figure(d, "fig3_gradient.png", 6.0, f"**Fig. 3.** Prevalence of food insecurity by number of disadvantages (woman reference person, no schooling or incomplete primary education, income ≤1/4 minimum wage per capita, rural residence) and race, with 95% confidence intervals.")
    H2(d, "3.7 Intersectional MAIHDA")
    ma, mg = M["ia_total"], M["ia_grave"]
    ora, orb = coef("ia_total", "renda_catq1"), coef("ia_grave", "renda_catq1")
    rca, rcb = coef("ia_total", "raca_catnegra"), coef("ia_grave", "raca_catnegra")
    sxa, sxb = coef("ia_total", "sexo_catmulher"), coef("ia_grave", "sexo_catmulher")
    para(d, f"Strata carried {f1(ma.vpc_nulo*100)}% (95% CI {ci(ma.vpc_nulo_lo*100, ma.vpc_nulo_hi*100, 1)}) of the latent variance in any food insecurity and {f1(mg.vpc_nulo*100)}% ({ci(mg.vpc_nulo_lo*100, mg.vpc_nulo_hi*100, 1)}) in severe food insecurity (Table 5). Adding the five main effects reduced these to {f1(ma.vpc_B*100)}% and {f1(mg.vpc_B*100)}%, a PCV of {f1(ma.pcv*100)}% and {f1(mg.pcv*100)}%. Income dominated the fixed effects: the odds ratio for income ≤1/4 MW versus >2 MW was {est_ci(ora.oddsr, ora.or_lo, ora.or_hi, 1)} for any and {est_ci(orb.oddsr, orb.or_lo, orb.or_hi, 1)} for severe food insecurity, against {f2(rca.oddsr)} and {f2(rcb.oddsr)} for Black race and {f2(sxa.oddsr)} and {f2(sxb.oddsr)} for women (Supplementary Table S7).")
    para(d, f"Predicted prevalence ranged from {f1(lo_any.p_completo*100)}% (non-Black man, secondary or tertiary education, income above 2 MW, rural) to {f1(hi_any.p_completo*100)}% (Black woman, no schooling or incomplete primary education, income ≤1/4 MW, urban) for any food insecurity, and from {f1(lo_sev.p_completo*100)}% to {f1(hi_sev.p_completo*100)}% for severe food insecurity. Stratum residuals did not indicate systematic departures from additivity: {int(ma.n_z_gt196)} of 160 strata exceeded |z| = 1.96 for any food insecurity, as expected by chance, and none for severe food insecurity; the largest difference between full-model and additive predictions was {f1(ma.max_abs_dif_pp)} percentage points (median {f1(ma.med_abs_dif_pp)}) and {f1(mg.max_abs_dif_pp)} points (median {f2(mg.med_abs_dif_pp)}), respectively (Fig. 4). Discriminatory accuracy was moderate: the AUC of stratum predictions was {f2(AUC['ia_total']['p_completo'])} for any and {f2(AUC['ia_grave']['p_completo'])} for severe food insecurity, against {f2(AUC['ia_total']['renda'])} and {f2(AUC['ia_grave']['renda'])} for income alone. With approximate weights, VPC was {f1(ma.w_vpc_nulo*100)}% and {f1(mg.w_vpc_nulo*100)}% and PCV {f1(ma.w_pcv*100)}% and {f1(mg.w_pcv*100)}% (Supplementary Table S8).")
    table_5(d)
    figure(d, "fig4_maihda_strata.png", 6.2, "**Fig. 4.** Predicted prevalence of food insecurity in the 160 intersectional strata, ranked by predicted risk, from the full model (main effects plus stratum residual; dots) and from additive main effects alone (open circles).")


def build_discussion(d):
    p_ = PAIR("inc", "ia_grave"); pe = PAIR("edu", "ia_grave")
    ma, mg = M["ia_total"], M["ia_grave"]
    bk_w = PAIR("sex", "ia_total"); bk_ws = PAIR("sex", "ia_grave")
    q1, q5 = T6("negra", "renda", "q1", "ia_total"), T6("negra", "renda", "q5", "ia_total")
    s1, s5 = T6("negra", "renda", "q1", "ia_grave"), T6("negra", "renda", "q5", "ia_grave")
    bf = AT("Negra", "renda_e_demais", "ia_total")
    wsi = AT("Mulher", "so_renda", "ia_grave"); wi = AT("Mulher", "renda_e_demais", "ia_total")
    ruf, rsf = AT("Rural", "renda_e_demais", "ia_total"), AT("Rural", "renda_e_demais", "ia_grave")
    sx = [T6("mulher", "renda", q, "ia_total") for q in ["q1", "q2", "q3", "q4", "q5"]]
    H2(d, "4.1 Principal findings")
    para(d, f"In a national survey of {n0(NA)} Brazilian households, intersectional inequality in food insecurity was large: predicted prevalence ranged from {f1(lo_any.p_completo*100)}% to {f1(hi_any.p_completo*100)}% across strata, and strata carried {f1(ma.vpc_nulo*100)}% to {f1(mg.vpc_nulo*100)}% of the latent variance. It was also, almost entirely, additive on the logit scale (PCV above 99%) and dominated by income. Black race retained an association with food insecurity after accounting for income and education, sex did so less, and rural residence turned protective. Pairwise interactions were sub-multiplicative, with super-additive excess risk concentrated in severe food insecurity, and no three-way interaction was detected.")
    H2(d, "4.2 Additive, multiplicative and multilevel views")
    para(d, f"These results are compatible once the scale is made explicit. When two exposures each carry a large risk ratio, the joint risk exceeds the sum of the separate risks even if it falls short of their product; the expected RERI under multiplicativity for race and income in severe food insecurity was {f2(p_.reri_esperado_nulo_mult)}, and we observed {f2(p_.reri)}. A positive RERI here therefore reflects the arithmetic of large relative risks on a low baseline and is smaller than multiplicativity would predict, so it should not be read as synergy or as evidence of a distinct mechanism {P('vanderweele2014','bauer2014')}. Because MAIHDA on the logit scale treats additive main effects as multiplicative on the odds, a PCV near 100% agrees with ratios of risk ratios close to, though for education and income below, one. On the risk scale, departure from additivity is real and not negligible for severe food insecurity, where it accounted for {f1(jd('inc','ia_grave')[3]/jd('inc','ia_grave')[0]*100)}% of the joint disparity for race and income and {f1(jd('edu','ia_grave')[3]/jd('edu','ia_grave')[0]*100)}% for race and education, against {f1(jd('inc','ia_total')[3]/jd('inc','ia_total')[0]*100)}% for race and income in any food insecurity; it is nonetheless smaller than a purely multiplicative process would generate. The additive-scale excess remains meaningful as a measure of burden: {f1(p_.p11)}% of Black households with income at or below 1/4 MW reported severe food insecurity, against {f1(p_.p00)}% of non-Black households above that income. The dependence of interaction on scale is documented empirically: among US college women, race and sexual orientation showed no additive-scale interaction but a negative multiplicative one {P('reynolds2021')}. Our findings echo a longitudinal MAIHDA of mental health in which inequalities were partly multiplicative but mostly additive {P('bell2024')}, and they show why reporting both scales, as recommended {P('knol2012')}, prevents contradictory conclusions from one dataset. The stratum residuals, few and small, are consistent with the tutorial's observation that VPC falls to near zero and PCV rises to near 100% once main effects enter {P('evans2024a')}. The additive share is outcome-specific: interaction-only variance was 1.1% in Swedish COPD {P('axelsson2018')}, whereas main effects explained 46% for obesity in Brazil {P('fanton2026')}, so a PCV near 100% for food insecurity is an empirical finding about a strongly income-driven outcome, not a property of the method.")
    ji = {dd: jd("inc", dd) for dd in ["ia_total", "ia_grave"]}
    para(d, f"Statistical interaction is neither necessary nor sufficient for intersectional processes. The theory concerns interlocking systems of power that position people differently, not the form of a regression term {P('bauer2014','bowleg2012')}; a Black woman with low income can face compounded exposure to racism, sexism and poverty even if the effects on an outcome combine roughly multiplicatively. Where interaction is present it is best read descriptively, as part of a joint disparity, {P('jackson2016')}: for race and income, interaction accounted for {f1(ji['ia_total'][3]/ji['ia_total'][0]*100)}% of the joint disparity in any and {f1(ji['ia_grave'][3]/ji['ia_grave'][0]*100)}% in severe food insecurity, the rest being the separate contributions of race and income. That an additive framework fits does not make the disparity small: the additive terms are themselves products of structural processes.")
    H2(d, "4.3 Race, income and education")
    para(d, f"Income explained about half of the crude racial association, in line with evidence that socioeconomic position accounts for part but not all of racial health gaps {P('farmer2005')} and that racism shapes both socioeconomic position and health {P('bailey2017','williams2019')}, but an association of {f2(bf.rr)} for any food insecurity persisted. Its pattern across strata is informative. The relative gap was larger at higher income (adjusted risk ratio {f2(q1.rr_ajustado)} to {f2(q5.rr_ajustado)} for any, and {f2(s1.rr_ajustado)} to {f2(s5.rr_ajustado)} for severe food insecurity) and education, whereas the absolute gap was larger at lower income ({f1(q1.dif_pp)} to {f1(q5.dif_pp)} percentage points). Both are true and answer different questions. Relative gaps are mechanically larger when baseline risk is small, so the upward slope should not be overstated, but it is consistent with diminished returns, whereby Black households obtain smaller protection from income and education {P('farmer2005','assari2018','bakhtiari2022')}, and with the finding in Salvador that Black women's households remained at higher odds in socioeconomically favourable strata {P('silva2022a')}. Evidence specific to food insecurity points the same way: Black households in Canada had higher predicted risk across nearly all sources of household income and relatively homogeneous risk across education {P('dhunna2021')}; in the United States disparities varied across education, poverty and home ownership but persisted within them {P('berning2024')}; and in Rio de Janeiro they persisted at higher income {P('camara2026')}. Greater racial inequality among those with higher socioeconomic status has been reported for self-rated health {P('brown2016')}. Fundamental-cause theory offers an interpretation {P('phelanlink2015')}, and the racism scholarship for Brazil identifies institutional and structural routes {P('werneck2016','costa2026')}, but a cross-sectional survey cannot identify mechanisms. Part of the residual could also reflect income measured as current household per capita income, which omits wealth, volatility and debt. Earlier estimates are not directly comparable, since they contrast profiles that also combine marital status and children, but {narr('luiz2026')} likewise found that racial inequalities persisted in 2023 even as food insecurity fell.")
    H2(d, "4.4 Sex and residence")
    para(d, f"The sex association was smaller and largely income-dependent: for severe food insecurity, income adjustment left a ratio of {est_ci(wsi.rr, wsi.ic95_low, wsi.ic95_high)}. For any food insecurity it persisted across income bands ({f2(min(x.rr_ajustado for x in sx))} to {f2(max(x.rr_ajustado for x in sx))}), consistent with evidence that economic factors explain much but not all of the gender gap {P('broussard2019','ivers2011')}. The reference person is not the only adult, and women-headed households differ in composition, notably single mothers with children, which earlier Brazilian studies found to raise risk {P('santos2023')}. We did not model marital status or children, so the sex coefficient describes women-headed households and not women as individuals, and headship is an imperfect targeting criterion {P('buvinic1997')}. The race-by-sex interaction was small (ratio of risk ratios {f2(bk_w.ror)} for any food insecurity; not significant for severe), and Black women's households had the highest joint risk ({est_cib(bk_w.rr11, bk_w.rr11_lo, bk_w.rr11_hi)} for any and {est_cib(bk_ws.rr11, bk_ws.rr11_lo, bk_ws.rr11_hi)} for severe food insecurity relative to non-Black men's), as expected from two independent associations. This differs in emphasis from the very high odds in Belo Horizonte, where the socioeconomic index, reference groups and outcome definition differ {P('ferreira2025')}. The rural reversal (risk ratio {f2(ruf.rr)} and {f2(rsf.rr)} after adjustment) shows that rural excess risk in Brazil is a poverty effect; it should not be read as rural life being protective, since own production and informal food access may not be captured by income. The EBIA shows invariance between urban and rural areas {P('cezimbra2026b')}. The regional analysis of race × sex hints at geographic heterogeneity, but intervals were wide and a study with more power per region is needed.")
    H2(d, "4.5 Implications")
    para(d, f"The MAIHDA results carry a caution for practice. Strata explained much variance, but discriminatory accuracy was moderate (AUC {f2(AUC['ia_total']['p_completo'])} and {f2(AUC['ia_grave']['p_completo'])}) and barely above that of income alone, so most variation lies within strata, as in the original MAIHDA application {P('evans2018')}. Targeting by race-sex-education profiles would be expected to misclassify many households, and income remains the strongest single marker. For policy, the dominance of income supports income-based social protection as the main lever on absolute inequality, while the persistence of the racial gap at every income level, and its larger relative size at higher income, suggests that income policies alone would not close it. Our data cannot say which complementary measures would work, and the tension between universal and targeted strategies {P('frohlich2008')} deserves attention when evaluating them.")
    para(d, f"Future work should use individual-level measures to separate household from individual gender effects; add wealth, income volatility, marital status and children; follow households over time, extending MAIHDA to trajectories {P('bell2024')}; develop survey-weighted MAIHDA; and test measurement invariance within intersectional strata rather than by single characteristics.")
    H2(d, "4.6 Strengths and limitations")
    para(d, f"Strengths include a large national sample, a single validated outcome measure, both interaction scales and benchmarks for each, complementary methods, and fully reproducible analyses. Limitations should be considered. First, the design is cross-sectional and descriptive; associations are not causal effects. Second, exposures describe the reference person, whereas food insecurity is measured for the household, so we cannot address intrahousehold allocation, and the reference person may not be the person who provisions food. Third, sex is binary and gender identity is not measured, and race is dichotomised, hiding differences between pretos and pardos and excluding Indigenous and Asian people from separate analysis. Fourth, income is a current per capita band; we lack wealth, income volatility, marital status, children, age and region, some of which matter {P('santos2023')}. Fifth, the MAIHDA models are unweighted, as in the standard procedure; approximate weights left VPC and PCV essentially unchanged, but the estimates describe the sample structure rather than population values, shrinkage implicitly weights strata equally {P('bashir2026')}, and results can depend on how strata are configured {P('humbert2024')}; the single-level results do not depend on these choices. Sixth, the EBIA shows approximate invariance across sex, race and education in 2013 data, but comparability across intersectional groups is not guaranteed {P('cezimbra2026a','bauer2019')}. Seventh, we ran many comparisons without adjustment and emphasise estimation over significance. Finally, the survey covers a single quarter, and the pattern may differ in other years, such as those with higher prevalence.")


def build_conclusions(d):
    para(d, "Intersectional inequality in household food insecurity in Brazil is large but is largely the accumulation of independent disadvantages, dominated by income. Racial disadvantage persists at every level of income and education, larger in relative terms at higher income and in absolute terms at lower income. Joint effects are super-additive but sub-multiplicative, so conclusions about synergy depend on the scale; on the risk scale, interaction accounts for a small share of the joint disparity in any food insecurity but about a fifth of it in severe food insecurity. Studies of intersectional inequality should report both scales and, where many strata are involved, the additive share of between-strata variance, and should distinguish excess risk as a measure of burden from evidence of mechanism.")


# ============================================================== tabelas
LAB1 = {"nao negra": "Non-Black", "negra (preta+parda)": "Black (preta or parda)", "homem": "Man", "mulher": "Woman",
        "sem_fund": "No schooling or incomplete primary", "fund_med": "Complete primary or incomplete secondary",
        "medio_comp": "Complete secondary", "superior": "Tertiary", "q1": "≤1/4", "q2": ">1/4–1/2", "q3": ">1/2–1", "q4": ">1–2",
        "q5": ">2", "urbano": "Urban", "rural": "Rural"}


def table_1(d):
    rows = []; groups = []
    def add(head, var, cats):
        groups.append(len(rows)); rows.append([head, "", "", "", ""])
        for c in cats:
            r = T1(var, c)
            rows.append([LAB1[c], n0(r.n), f1(r.pct_pond_amostra),
                         f"{f1(r.ia_total_pct)} ({ci(r.ia_total_lo, r.ia_total_hi, 1)})",
                         f"{f1(r.ia_grave_pct)} ({ci(r.ia_grave_lo, r.ia_grave_hi, 1)})"])
    rows.append(["All households", n0(tot.n), "100.0", f"{f1(tot.ia_total_pct)} ({ci(tot.ia_total_lo, tot.ia_total_hi, 1)})",
                 f"{f1(tot.ia_grave_pct)} ({ci(tot.ia_grave_lo, tot.ia_grave_hi, 1)})"])
    add("Race/colour", "Cor/raca", ["nao negra", "negra (preta+parda)"])
    add("Sex", "Sexo", ["homem", "mulher"])
    add("Education", "Instrucao", ["sem_fund", "fund_med", "medio_comp", "superior"])
    add("Per capita income (minimum wages)", "Renda pc", ["q1", "q2", "q3", "q4", "q5"])
    add("Residence", "Situacao", ["urbano", "rural"])
    caption(d, "**Table 1.** Households, weighted distribution and prevalence of food insecurity by characteristic of the reference person, PNADC 2023.")
    table(d, ["Characteristic", "Households (n)", "Weighted %", "Any food insecurity, % (95% CI)", "Severe food insecurity, % (95% CI)"], rows,
          widths=[5.4, 2.2, 2.0, 3.4, 3.4], group_rows=groups,
          note="Weighted percentages and confidence intervals (Taylor linearisation, primary sampling unit). MW, minimum wage.")


def table_2(d):
    rows = []
    labs = [("Negra", "Black race"), ("Mulher", "Woman"), ("Sem instrução", "No schooling/incomplete primary"), ("Rural", "Rural residence")]
    for v, lab in labs:
        for dd, dl in [("ia_total", "Any"), ("ia_grave", "Severe")]:
            b, r1, r2 = AT(v, "bruto", dd), AT(v, "so_renda", dd), AT(v, "renda_e_demais", dd)
            rows.append([lab if dl == "Any" else "", dl,
                         est_ci_short(b.rr, b.ic95_low, b.ic95_high), est_ci_short(r1.rr, r1.ic95_low, r1.ic95_high),
                         est_ci_short(r2.rr, r2.ic95_low, r2.ic95_high), f"{r1.atenuacao_pct:.0f}", f"{r2.atenuacao_pct:.0f}"])
    caption(d, "**Table 2.** Risk ratios (95% CI) for food insecurity by characteristic: crude, adjusted for income, and adjusted for income and the other characteristics.")
    table(d, ["Characteristic", "Outcome", "Crude", "Adjusted for income", "Adjusted for income and others", "Attenuation, income (%)", "Attenuation, full (%)"], rows,
          widths=[3.6, 1.5, 2.6, 2.6, 2.7, 1.9, 1.9],
          note="Poisson regression with robust variance, survey weights, standard errors clustered on primary sampling units. Attenuation is the percentage reduction of the crude log risk ratio; values above 100 indicate reversal of direction.")


def table_3(d):
    rows = []
    for k, lab in [("sex", "Race × sex (woman)"), ("edu", "Race × education (≤ incomplete primary)"), ("inc", "Race × income (≤1/4 MW)"), ("res", "Race × residence (rural)")]:
        for dd, dl in [("ia_total", "Any"), ("ia_grave", "Severe")]:
            r = PAIR(k, dd)
            rows.append([lab if dl == "Any" else "", dl, est_ci_short(r.rr10, r.rr10_lo, r.rr10_hi), est_ci_short(r.rr01, r.rr01_lo, r.rr01_hi),
                         est_ci_short(r.rr11, r.rr11_lo, r.rr11_hi), est_ci_short(r.reri, r.reri_lo_delta, r.reri_hi_delta),
                         f2(r.reri_esperado_nulo_mult), est_ci_short(r.ror, r.ror_lo, r.ror_hi)])
    caption(d, "**Table 3.** Joint risk ratios and interaction between race and four characteristics on additive and multiplicative scales.")
    table(d, ["Pair", "Outcome", "Black only, RR", "Other only, RR", "Both, RR", "RERI (95% CI)", "RERI expected if multiplicative", "Ratio of RRs (95% CI)"], rows,
          widths=[3.4, 1.4, 2.1, 2.1, 2.2, 2.4, 1.8, 2.3],
          note="Reference: non-Black without the second characteristic. RERI, relative excess risk due to interaction (delta-method interval); ratio of RRs = RR(both)/[RR(Black only) × RR(other only)]. MW, minimum wage.")


def table_4(d):
    rows = []
    labs = [("negra", "Black race"), ("mulher", "Woman"), ("sem_fund", "No schooling/incomplete primary"), ("renda_q1", "Income ≤1/4 MW"), ("rural", "Rural residence"),
            ("negra:mulher", "Black × woman"), ("negra:sem_fund", "Black × no schooling/incomplete primary"), ("negra:renda_q1", "Black × income ≤1/4 MW"), ("negra:rural", "Black × rural")]
    for t, lab in labs:
        a, s = T3(t, "ia_total"), T3(t, "ia_grave")
        rows.append([lab, est_ci_short(a.rr, a.lo, a.hi), pfmt(a.p), est_ci_short(s.rr, s.lo, s.hi), pfmt(s.p)])
    caption(d, "**Table 4.** Mutually adjusted model with main effects and race-by-characteristic products: risk ratios (95% CI).")
    table(d, ["Term", "Any food insecurity", "p", "Severe food insecurity", "p"], rows, widths=[6.0, 3.6, 1.4, 3.6, 1.4], group_rows=(),
          note="Product terms are ratios of risk ratios (multiplicative interaction). MW, minimum wage.")


def table_5(d):
    ma, mg = M["ia_total"], M["ia_grave"]
    A, G = AUC["ia_total"], AUC["ia_grave"]
    rows = [
        ["Between-stratum variance, null model (σ²A)", f3(ma.sigma2_nulo), f3(mg.sigma2_nulo)],
        ["VPC, null model, % (95% CI)", f"{f1(ma.vpc_nulo*100)} ({ci(ma.vpc_nulo_lo*100, ma.vpc_nulo_hi*100, 1)})", f"{f1(mg.vpc_nulo*100)} ({ci(mg.vpc_nulo_lo*100, mg.vpc_nulo_hi*100, 1)})"],
        ["Between-stratum variance, main-effects model (σ²B)", f3(ma.sigma2_B), f3(mg.sigma2_B)],
        ["VPC, main-effects model, %", f1(ma.vpc_B * 100), f1(mg.vpc_B * 100)],
        ["PCV, %", f1(ma.pcv * 100), f1(mg.pcv * 100)],
        ["AUC, stratum predictions (full model)", f2(A["p_completo"]), f2(G["p_completo"])],
        ["AUC, additive main effects only", f2(A["p_aditivo"]), f2(G["p_aditivo"])],
        ["AUC, income band only", f2(A["renda"]), f2(G["renda"])],
        ["Strata with |z| > 1.96 (of 160; about 8 expected by chance)", str(int(ma.n_z_gt196)), str(int(mg.n_z_gt196))],
        ["Largest full-minus-additive difference, percentage points", f1(ma.max_abs_dif_pp), f1(mg.max_abs_dif_pp)],
        ["Median absolute difference, percentage points", f2(ma.med_abs_dif_pp), f2(mg.med_abs_dif_pp)],
        ["Approximate weights: VPC null / PCV, %", f"{f1(ma.w_vpc_nulo*100)} / {f1(ma.w_pcv*100)}", f"{f1(mg.w_vpc_nulo*100)} / {f1(mg.w_pcv*100)}"],
    ]
    caption(d, "**Table 5.** Intersectional MAIHDA across 160 strata: variance partition, proportional change in variance, discriminatory accuracy and departures from additivity.")
    table(d, ["Measure", "Any food insecurity", "Severe food insecurity"], rows, widths=[8.6, 3.6, 3.6],
          note="Logistic models with a stratum random intercept; VPC on the latent scale. PCV, proportional change in stratum variance; AUC, area under the ROC curve of stratum-predicted probabilities. Models unweighted except the last row.")


# ============================================================== blocos auxiliares
HIGHLIGHTS = [
    "Household food insecurity analysed across 160 intersectional strata in Brazil",
    "Income explains about half of the racial gap; the rest persists at all income levels",
    "Joint effects are super-additive but sub-multiplicative: the scale changes the story",
    "Additive main effects explain over 99% of between-strata variance (MAIHDA)",
    "Relative racial gaps widen with income while absolute gaps narrow",
]

AI_STATEMENT = ("During the preparation of this work the authors used Claude (Anthropic) to assist with literature searching, drafting and translating "
                "the manuscript, and writing analysis and figure code. After using this tool, the authors reviewed and edited the content and take full "
                "responsibility for the content of the publication.")


def build_references(d):
    for k in sorted(USED, key=lambda k: (label(k).lower(), k)):
        p = d.add_paragraph(); rich(p, ref_entry(k), size=10); p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.9)
        WC["refs"] += wcount(ref_entry(k))


def main():
    # ---------- manuscrito cego
    d = new_doc(line_numbers=True, spacing=1.5)
    p = d.add_paragraph(); r = p.add_run("Intersectional inequalities in household food insecurity in Brazil: additive, multiplicative and multilevel evidence from a national survey, 2023")
    r.bold = True; r.font.size = Pt(15); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    H1(d, "Abstract")
    ab = build_abstract(); para(d, ab, bucket="abstract")
    p = d.add_paragraph(); rich(p, "*Keywords:* food insecurity; intersectionality; multilevel analysis of individual heterogeneity and discriminatory accuracy (MAIHDA); racial inequality; interaction; Brazil")
    WC["abstract"] += 16
    H1(d, "1. Introduction"); build_intro(d)
    H1(d, "2. Methods"); build_methods(d)
    H1(d, "3. Results"); build_results(d)
    H1(d, "4. Discussion"); build_discussion(d)
    H1(d, "5. Conclusions"); build_conclusions(d)
    H1(d, "Declaration of generative AI and AI-assisted technologies in the writing process")
    para(d, AI_STATEMENT, bucket="decl")
    H1(d, "Data availability")
    para(d, "The survey microdata are public and available from IBGE (https://www.ibge.gov.br). Scripts that reproduce every table and figure, and the derived tables, will be deposited in an open repository on acceptance [repository and DOI to be added].", bucket="decl")
    H1(d, "References")
    build_references(d)
    total = sum(WC.values())
    d.save(OUT / "02_Manuscript_blinded.docx")
    print("PALAVRAS (SSM conta tudo, limite 9.000):", WC, "TOTAL =", total)
    unused = [k for k in DOIS if k not in USED]
    print("refs usadas:", len(USED), "| nao usadas:", unused)
    hl = [(h, len(h)) for h in HIGHLIGHTS]
    print("highlights (limite 85 caracteres):", hl)
    return ab


if __name__ == "__main__":
    main()
