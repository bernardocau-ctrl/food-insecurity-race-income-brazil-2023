# -*- coding: utf-8 -*-
"""
auditoria_submissao_ije_20260926.py
ETAPA 5 (submissao): contagem de palavras sob varias regras, referencias cruzadas a tabelas/figuras (incl. suplementares) apos a
renumeracao, conteudo numerico das tabelas suplementares que repetem estimativas do corpo (S2, S4, S5), arquivos de figura,
coerencia entre o arquivo de tabelas e o manuscrito. NAO estima nada.
"""
import re, io, hashlib
from pathlib import Path
import pandas as pd
from docx import Document
from PIL import Image

BASE = Path(__file__).resolve().parent.parent
OUT = BASE / "MANUSCRITO_IJE_20260925"; AUD = OUT / "auditoria"; V2 = BASE / "dados" / "v3_rodadaD_C"
A = Document(OUT / "A_Manuscript_IJE_rev4.docx"); B = Document(OUT / "B_Supplementary_material_IJE_rev4.docx"); TB = Document(OUT / "Tables_1-3_IJE_rev4.docx")
log = []; L = lambda *a: log.append(" ".join(str(x) for x in a))

# ---------------- (a) contagem de palavras
paras = [p.text for p in A.paragraphs]
def idx(t, start=0): return next(i for i in range(start, len(paras)) if paras[i].strip() == t)
i_intro, i_decl = idx("Introduction"), idx("Declarations")
CAP = re.compile(r"^(Table \d|Figure \d|Weighted percentages|Crude, survey-weighted estimates \(no covariates\))")
body = [t for t in paras[i_intro:i_decl] if t.strip()]
heads = {"Introduction", "Methods", "Results", "Discussion", "Conclusion", "Design and data", "Outcome", "Exposures", "Statistical analysis", "Sample and prevalence", "Race/colour and income", "Race/colour and education", "Relative and absolute racial gaps across income", "Complementary MAIHDA"}
txt = [t for t in body if not CAP.match(t)]
words = lambda t: len(t.split())
w_all = sum(words(t) for t in txt); w_nohead = sum(words(t) for t in txt if t.strip() not in heads)
cite = re.compile(r"\[\d+(?:[,\-–]\s*\d+)*\]")
w_nocite = sum(words(cite.sub("", t)) for t in txt if t.strip() not in heads)
w_alpha = sum(len(re.findall(r"[A-Za-zÀ-ÿ]+(?:['’\-][A-Za-zÀ-ÿ]+)*", cite.sub("", t))) for t in txt if t.strip() not in heads)
i_abs, i_km = idx("Abstract"), idx("Key Messages")
abs_w = sum(words(t) for t in paras[i_abs:i_km] if re.match(r"(Background|Methods|Results|Conclusions):", t))
R = {"texto_principal_com_titulos_(split)": w_all, "texto_principal_sem_titulos_(split)": w_nohead, "sem_titulos_e_sem_citacoes_[n]": w_nocite, "apenas_palavras_alfabeticas_(sem_numeros/citacoes)": w_alpha, "resumo_(split)": abs_w}
L("## Contagem de palavras (recontada do .docx rev4)"); [L(f"- {k}: {v}") for k, v in R.items()]

# ---------------- (b) referencias cruzadas
capA = {}
for p in A.paragraphs:
    m = re.match(r"^(Table \d|Figure \d)\.", p.text)
    if m: capA[m.group(1)] = p.text
capB = {}
for p in B.paragraphs:
    m = re.match(r"^Supplementary (Table|Figure) (S\d+)([ab])?\.", p.text)
    if m: capB[f"{m.group(1)} {m.group(2)}{m.group(3) or ''}"] = p.text
capB_base = {}
for k in capB:
    capB_base.setdefault(re.sub(r"[ab]$", "", k), []).append(k)
def all_text(doc):
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells: out.append(c.text)
    return out
cites = []
for name, doc in (("A", A), ("B", B)):
    for t in all_text(doc):
        if name == "A" and t.strip().startswith("Supplementary"): pass
        for m in re.finditer(r"Supplementary (Tables?|Figures?) ((?:S\d+[ab]?(?:(?:, | and | to |-)S?\d+[ab]?)*))", t):
            kind = "Table" if m.group(1).startswith("Table") else "Figure"
            toks = re.split(r", | and | to |-", m.group(2)); nums = []
            if "-" in m.group(2) and len(toks) == 2:
                a, b = [int(re.sub(r"\D", "", x)) for x in toks]; nums = [f"S{i}" for i in range(a, b + 1)]
            else: nums = [x if x.startswith("S") else "S" + x for x in toks]
            for n in nums: cites.append((name, kind, n, t[max(0, m.start() - 60): m.end() + 30].replace("\n", " ")))
miss = [(n, k, s, c) for n, k, s, c in cites if f"{k} {s}" not in capB and f"{k} {s}" not in capB_base and f"{k} {s}a" not in capB]
L("\n## Referencias cruzadas a itens suplementares"); L(f"- Legendas encontradas no suplemento: tabelas {sorted(k for k in capB if k.startswith('Table'))}; figuras {sorted(k for k in capB if k.startswith('Figure'))}")
L(f"- Citacoes analisadas: {len(cites)}; sem legenda correspondente: {len(miss)}")
for m in miss: L("  - NAO RESOLVIDA:", m)
# tabelas do corpo citadas
citedT = sorted(set(re.findall(r"\bTables? (\d)(?:(?:-| and )(\d))?", " ".join(txt))))
L(f"- 'Table n' citadas no texto principal: {citedT}; legendas no manuscrito: {sorted(k for k in capA)}")
rows = []
for n, k, s, c in cites:
    key = f"{k} {s}"; cap = capB.get(key) or capB.get(key + "a") or ""
    rows.append({"em": n, "item": key, "contexto_da_citacao": c, "legenda_no_suplemento": cap[:200]})
pd.DataFrame(rows).drop_duplicates(["em", "item", "contexto_da_citacao"]).to_csv(AUD / "referencias_cruzadas_suplemento_rev4.csv", index=False, encoding="utf-8-sig")
# S1b -> tabelas citadas existem? (ja coberto acima)

# ---------------- (c) figuras
L("\n## Arquivos de figura")
for f in sorted((OUT / "figures").glob("*.png")):
    im = Image.open(f); L(f"- {f.name}: {im.size[0]}x{im.size[1]} px; dpi gravado {im.info.get('dpi')}; largura a 300 dpi = {im.size[0]/300*25.4:.0f} mm")
shp = A.inline_shapes[0]; rid = shp._inline.graphic.graphicData.pic.blipFill.blip.embed; blob = A.part.related_parts[rid].blob
L(f"- Figura 1 embutida no manuscrito == arquivo separado fig1_interaction_scales.png: {hashlib.md5(blob).hexdigest() == hashlib.md5((OUT / 'figures' / 'fig1_interaction_scales.png').read_bytes()).hexdigest()}")
alt = shp._inline.docPr.get("descr") or ""
capfig = capA.get("Figure 1", "")
altpar = [p for p in paras if p.startswith("Figure 1, alt text:")]
L(f"- Alt text embutido na imagem: {bool(alt)}; alt text tambem como paragrafo no documento principal: {bool(altpar)}; legenda no documento: {bool(capfig)}")
(OUT / "figures" / "Figure1_legend_and_alt_text.txt").write_text("FIGURE 1. LEGEND (paste into the 'Caption/Legend' box; identical to the main document)\n\n" + capfig + "\n\nALT TEXT\n\n" + alt + "\n", encoding="utf-8")

# ---------------- (d) coerencia arquivo de tabelas x manuscrito
def tbl_txt(doc): return [[c.text for c in r.cells] for t in doc.tables for r in t.rows]
L("\n## Arquivo separado de tabelas"); L(f"- Tabelas no manuscrito: {len(A.tables)}; no arquivo Tables_1-3_IJE_rev4.docx: {len(TB.tables)}; conteudo das celulas identico: {tbl_txt(A) == tbl_txt(TB)}")

# ---------------- (e) suplemento: S2, S4, S5 contra os CSV
t1 = pd.read_csv(V2 / "t1_descritiva.csv"); t2 = pd.read_csv(V2 / "t2_pares.csv")
NB = " "
def tok(cell): return [float(x.replace("−", "-").replace(NB, "").replace(" ", "")) for x in re.findall(r"(?:(?<![\d)])[−-])?\d[\d  ]*\.?\d*", cell)]
def ok(cell, vals, dec):
    g = tok(cell); return len(g) == len(vals) and all(abs(a - round(v, dec)) < 0.5 * 10 ** -dec + 1e-12 for a, v in zip(g, vals))
sup_tables = B.tables; res = []
LAB = {"Non-Black": ("Cor/raca", "nao negra"), "Black (preta or parda)": ("Cor/raca", "negra (preta+parda)"), "Man": ("Sexo", "homem"), "Woman": ("Sexo", "mulher"),
       "No schooling or incomplete primary": ("Instrucao", "sem_fund"), "Complete primary or incomplete secondary": ("Instrucao", "fund_med"), "Complete secondary": ("Instrucao", "medio_comp"),
       "Tertiary (complete or incomplete)": ("Instrucao", "superior"), "≤1/4": ("Renda pc", "q1"), ">1/4-1/2": ("Renda pc", "q2"), ">1/2-1": ("Renda pc", "q3"), ">1-2": ("Renda pc", "q4"), ">2": ("Renda pc", "q5"),
       "Urban": ("Situacao", "urbano"), "Rural": ("Situacao", "rural"), "All households": ("Total", "Brasil")}
# localizar tabelas pelo cabecalho
def find_tbl(first_hdr, second=None):
    for t in sup_tables:
        h = [c.text for c in t.rows[0].cells]
        if h[0] == first_hdr and (second is None or h[1] == second): return t
S2 = find_tbl("Characteristic", "Households (n)"); n2 = 0
for r in S2.rows[1:]:
    c = [x.text for x in r.cells]
    if c[0] in LAB:
        row = t1[(t1.variavel == LAB[c[0]][0]) & (t1.categoria == LAB[c[0]][1])].iloc[0]
        res.append(("S2 " + c[0], ok(c[1], [row["n"]], 0) and ok(c[2], [row["pct_pond_amostra"]], 1) and ok(c[3], [row.ia_total_pct, row.ia_total_lo, row.ia_total_hi], 1) and ok(c[4], [row.ia_grave_pct, row.ia_grave_lo, row.ia_grave_hi], 1)))
S4 = find_tbl("Pair", "Outcome")
PARMAP = {"Race/colour × income (≤1/4 MW)": "Race x Income (<= 1/4 MW per capita)", "Race/colour × education (≤ incomplete primary)": "Race x Education (<= incomplete primary)", "Race/colour × sex (woman)": "Race x Sex (woman)", "Race/colour × residence (rural)": "Race x Residence (rural)"}
cur = None
for r in S4.rows[1:]:
    c = [x.text for x in r.cells]; cur = PARMAP.get(c[0], cur); y = "ia_total" if c[1] == "Any" else "ia_grave"; q = t2[(t2.par == cur) & (t2.desfecho == y)].iloc[0]
    res.append((f"S4 {cur[:22]}/{y}", ok(c[2], [q.rr10, q.rr10_lo, q.rr10_hi], 2) and ok(c[3], [q.rr01, q.rr01_lo, q.rr01_hi], 2) and ok(c[4], [q.rr11, q.rr11_lo, q.rr11_hi], 2) and ok(c[5], [q.reri, q.reri_lo_delta, q.reri_hi_delta], 2) and ok(c[6], [q.reri_esperado_nulo_mult], 2) and ok(c[7], [q.ror, q.ror_lo, q.ror_hi], 2)))
S5 = None
for t in sup_tables:
    if t.rows[0].cells[0].text == "Pair" and "Prevalence" in t.rows[0].cells[2].text: S5 = t
cur = None
for r in S5.rows[1:]:
    c = [x.text for x in r.cells]; cur = PARMAP.get(c[0], cur); y = "ia_total" if c[1] == "Any" else "ia_grave"; q = t2[(t2.par == cur) & (t2.desfecho == y)].iloc[0]
    okp = ok(c[2], [q.p00, q.p10, q.p01, q.p11], 1) and ok(c[3], [q.n00, q.n10, q.n01, q.n11], 0) and ok(c[4], [q.reri, q.reri_lo_delta, q.reri_hi_delta], 2)
    j = q.p11 - q.p00; a = q.p10 - q.p00; b = q.p01 - q.p00; it = q.p11 - q.p10 - q.p01 + q.p00
    okd = ok(c[6], [j, a, b, it], 1) and ok(c[7], [it / j * 100], 1)
    res.append((f"S5 {cur[:22]}/{y}", okp and okd))
L("\n## Tabelas suplementares que repetem estimativas do corpo (S2, S4, S5) contra os CSV"); L(f"- {sum(1 for _, o in res if o)} de {len(res)} linhas conferem; falhas: {[n for n, o in res if not o]}")
(AUD / "auditoria_submissao_resumo_rev4.md").write_text("\n".join(log), encoding="utf-8"); print("\n".join(log))
